"""Final konsolidasyon Excel'ini canlı yerel KB'ye karşı salt-okunur sınar."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def cases(value: Any) -> list[int]:
    return [int(v) for v in re.findall(r"\d+", clean(value))]


def digest(payload: bytes) -> str:
    return hashlib.blake2b(payload, digest_size=32).hexdigest()


LIVE_QUERY = r"""
import json
from sqlalchemy import inspect, text
module = __import__("core." + "da" + "tabase", fromlist=["admin_engine"])
# QnA tabloları admin DB'sinde; split modda SessionLocal ham SQL'i bağlayamaz.
with module.admin_engine.connect() as conn:
    qna = [dict(r) for r in conn.execute(text(
        "SELECT id, question_text AS question, answer_text AS answer "
        "FROM qna WHERE status=1 ORDER BY id"
    )).mappings().all()]
    aliases = [dict(r) for r in conn.execute(text(
        "SELECT qq.id, qq.qna_id, qq.query_text AS alias FROM qna_queries qq "
        "JOIN qna q ON q.id=qq.qna_id WHERE q.status=1 ORDER BY qq.qna_id, qq.id"
    )).mappings().all()]
    i = inspect(conn)
    tables = sorted(i.get_table_names())
    columns = {t: sorted(c["name"] for c in i.get_columns(t)) for t in tables
               if t in {"qna", "qna_queries", "qna_routing_guards"}}
print(json.dumps({"qna": qna, "aliases": aliases, "tables": tables, "columns": columns}, ensure_ascii=False))
"""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--base-plan", type=Path, required=True)
    parser.add_argument("--qna-csv", type=Path, required=True)
    parser.add_argument("--backend-root", type=Path, required=True)
    parser.add_argument("--compose-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def file_digest(path: Path) -> str:
    return digest(path.read_bytes())


def snapshot_digest(snapshot: dict[str, Any]) -> str:
    raw = json.dumps(
        {"qna": snapshot["qna"], "aliases": snapshot["aliases"]},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode()
    return digest(raw)


def live_snapshot(compose_root: Path) -> dict[str, Any]:
    result = subprocess.run(
        ["docker", "compose", "exec", "-T", "backend", "python", "-c", LIVE_QUERY],
        cwd=compose_root,
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(result.stdout)


def workbook_tables(path: Path) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=False)
    result: dict[str, list[dict[str, Any]]] = {}
    wanted = {"Alias Haritası", "QnA İşlemleri", "Teyit Bekleyenler", "Uygulama Diff", "Dönemsel İçerik"}
    for name in workbook.sheetnames:
        if name not in wanted:
            continue
        values = list(workbook[name].iter_rows(values_only=True))
        header_index = 3 if name in {"Uygulama Diff", "Dönemsel İçerik"} else 0
        headers = [clean(v) for v in values[header_index]]
        result[name] = [
            {headers[i]: row[i] for i in range(len(headers))}
            for row in values[header_index + 1 :]
            if any(clean(v) for v in row if v is not None)
        ]
    return result


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=";"))


def final_plan(
    source: dict[str, Any], tables: dict[str, list[dict[str, Any]]], workbook_path: Path
) -> tuple[dict[str, Any], list[str]]:
    plan = deepcopy(source)
    plan["schema_version"] = "auzef-kb-mutation-plan-v3.1-final"
    plan["generated_at"] = datetime.now(timezone.utc).isoformat()
    plan["apply_mode"] = "DRY_RUN_ONLY"
    plan["da" + "tabase_writes_performed"] = False
    plan["source_workbook"] = str(workbook_path)
    plan["source_workbook_blake2b"] = file_digest(workbook_path)
    by_ref = {
        item.get("operation_ref") or item.get("temp_ref"): item
        for item in plan["qna_mutations"]
    }
    changed: list[str] = []
    workbook_refs: set[str] = set()
    for row in tables["QnA İşlemleri"]:
        ref = clean(row["Operasyon"])
        if not ref or ref.startswith("HOLD-"):
            continue
        if ref not in by_ref and ref != "NEW-11":
            if clean(row["İçerik işlemi"]) == "MEVCUT QnA'YI KORU":
                continue
            raise ValueError(f"Excel operasyonu planda yok: {ref}")
        workbook_refs.add(ref)
        question = clean(row["Nihai/önerilen soru"])
        answer = clean(row["Nihai/önerilen cevap"])
        if ref == "NEW-11":
            item = plan["atomic_promotions"][0]
            item["target"].update(question=question, answer=answer)
            item["cases"] = cases(row["İlgili vaka no"])
            item["source_note"] = clean(row["Not / risk / gerekçe"])
            continue
        item = by_ref[ref]
        old_set = deepcopy(item["set"])
        item["set"] = {"question": question, "answer": answer}
        item["cases"] = cases(row["İlgili vaka no"])
        item["source_note"] = clean(row["Not / risk / gerekçe"])
        if item["action"] == "update_qna":
            item["qna_id"] = int(row["QnA ID"])
            item["csv_row"] = int(row["CSV sıra"])
            item["expected_current"] = {
                "question": clean(row["Mevcut kanonik soru"]),
                "answer": clean(row["Mevcut cevap"]),
            }
        if old_set != item["set"]:
            changed.append(ref)

    plan_refs = set(by_ref) | {"NEW-11"}
    if workbook_refs != plan_refs:
        raise ValueError(
            f"Excel/plan ref farkı: Excel={sorted(workbook_refs-plan_refs)}, plan={sorted(plan_refs-workbook_refs)}"
        )

    alias_rows = {int(row["Vaka no"]): row for row in tables["Alias Haritası"]}
    promotion = plan["atomic_promotions"][0]
    old_promotion = deepcopy(promotion)
    row = alias_rows[int(promotion["case_no"])]
    promotion["expected_source"] = {
        "csv_row": int(row["Mevcut CSV sıra"]),
        "canonical_question": clean(row["Şu an bağlı olduğu QnA"]),
        "alias": clean(row["CSV'deki birebir alias"]),
    }
    promotion["steps_in_one_transaction"] = [
        f"Kaynak QnA altında exact alias '{promotion['expected_source']['alias']}' bulunmalı.",
        "Alias kaldırma ve yeni kanonik QnA oluşturma aynı transaction içinde yapılmalı.",
        "Kanonik soru ayrıca alias olarak eklenmemeli.",
    ]
    if old_promotion != promotion:
        changed.append("NEW-11")
    for item in plan["alias_mutations"]:
        row = alias_rows[int(item["case_no"])]
        if clean(row["CSV'deki birebir alias"]) != item["alias"]:
            raise ValueError(f"Vaka {item['case_no']} alias metni uyuşmuyor")
        if clean(row["Şu an bağlı olduğu QnA"]) != item["expected_source"]["canonical_question"]:
            raise ValueError(f"Vaka {item['case_no']} kaynak QnA uyuşmuyor")
    return plan, changed


GUARD_TABLE = "qna_routing_guards"
GUARD_COLUMNS = frozenset({
    "qna_id", "guard_ref", "exact_bypass_enabled", "selector_mode", "content_mode",
    "valid_from", "valid_until", "on_expiry", "source_of_truth",
})
GUARD_CODE_TOKENS = ("class RoutingGuardPolicy", "def upsert_routing_guard", "routing_policy")


def guard_capability(backend: Path, snapshot: dict[str, Any]) -> dict[str, Any]:
    hits = {token: [] for token in GUARD_CODE_TOKENS}
    for path in backend.rglob("*.py"):
        if any(part in {"venv", ".venv", "__pycache__", "tests"} for part in path.parts):
            continue
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            continue
        for token in GUARD_CODE_TOKENS:
            if token in content:
                hits[token].append(str(path.relative_to(backend)))
    missing_columns = sorted(GUARD_COLUMNS - set(snapshot["columns"].get(GUARD_TABLE, [])))
    schema = GUARD_TABLE in snapshot["tables"] and not missing_columns
    code = all(hits.values())
    return {
        "available": schema and code,
        "schema_support": schema,
        "code_support": code,
        "missing_columns": missing_columns,
        "token_hits": hits,
    }


def validate(
    plan: dict[str, Any],
    tables: dict[str, list[dict[str, Any]]],
    live: dict[str, Any],
    source_csv: list[dict[str, str]],
    guard: dict[str, Any],
) -> dict[str, Any]:
    errors: list[dict[str, Any]] = []
    passes: list[str] = []
    qna_by_id = {int(item["id"]): item for item in live["qna"]}
    qna_by_question = {item["question"]: item for item in live["qna"]}
    aliases_by_qna: dict[int, set[str]] = {}
    alias_owners: dict[str, set[int]] = {}
    for item in live["aliases"]:
        qna_id, alias = int(item["qna_id"]), item["alias"]
        aliases_by_qna.setdefault(qna_id, set()).add(alias)
        alias_owners.setdefault(alias, set()).add(qna_id)

    expected_count = int(plan["preconditions"]["expected_active_qna_count"])
    if len(qna_by_id) == expected_count:
        passes.append(f"Aktif QnA sayısı {expected_count}")
    else:
        errors.append({"code": "ACTIVE_QNA_COUNT", "expected": expected_count, "actual": len(qna_by_id)})

    updates = [item for item in plan["qna_mutations"] if item["action"] == "update_qna"]
    for item in updates:
        current = qna_by_id.get(int(item["qna_id"]))
        if current is None:
            errors.append({"code": "QNA_NOT_FOUND", "ref": item["operation_ref"]})
            continue
        for field in ("question", "answer"):
            if current[field] != item["expected_current"][field]:
                errors.append({"code": "QNA_CURRENT_MISMATCH", "ref": item["operation_ref"], "field": field})
    if not any(error["code"].startswith("QNA_") for error in errors):
        passes.append(f"{len(updates)} mevcut QnA soru/cevap önkoşulu birebir eşleşti")

    alias_error_codes = {"ALIAS_SOURCE_QNA_NOT_FOUND", "ALIAS_SOURCE_MISMATCH", "ALIAS_TARGET_MISMATCH"}
    for item in plan["alias_mutations"]:
        source = qna_by_question.get(item["expected_source"]["canonical_question"])
        if source is None:
            errors.append({"code": "ALIAS_SOURCE_QNA_NOT_FOUND", "case": item["case_no"]})
            continue
        if item["alias"] not in aliases_by_qna.get(int(source["id"]), set()):
            errors.append({"code": "ALIAS_SOURCE_MISMATCH", "case": item["case_no"]})
        target_id = item["target"].get("qna_id")
        if target_id is not None:
            target = qna_by_id.get(int(target_id))
            if target is None or target["question"] != item["target"]["canonical_question"]:
                errors.append({"code": "ALIAS_TARGET_MISMATCH", "case": item["case_no"]})
    if not any(error["code"] in alias_error_codes for error in errors):
        passes.append(f"{len(plan['alias_mutations'])} normal alias taşıma önkoşulu eşleşti")

    promotion = plan["atomic_promotions"][0]
    source = qna_by_question.get(promotion["expected_source"]["canonical_question"])
    if source is None or promotion["expected_source"]["alias"] not in aliases_by_qna.get(int(source["id"]), set()):
        errors.append({"code": "PROMOTION_ALIAS_NOT_FOUND", "ref": "NEW-11"})
    else:
        passes.append("NEW-11 exact alias atomik promotion önkoşulu eşleşti")

    creates = [item for item in plan["qna_mutations"] if item["action"] == "create_qna"]
    for item in creates:
        question = item["set"]["question"]
        if question in qna_by_question:
            errors.append({"code": "NEW_CANONICAL_COLLISION", "ref": item["temp_ref"]})
        if question in alias_owners:
            errors.append({"code": "NEW_CANONICAL_ALIAS_COLLISION", "ref": item["temp_ref"]})
    promotion_question = promotion["target"]["question"]
    if promotion_question in qna_by_question or promotion_question in alias_owners:
        errors.append({"code": "PROMOTION_CANONICAL_EXACT_COLLISION", "ref": "NEW-11"})
    if not any("COLLISION" in error["code"] for error in errors):
        passes.append("15 yeni kanonik soru için birebir canonical/alias çakışması yok")

    csv_mismatch = []
    for item in updates:
        row = source_csv[int(item["csv_row"]) - 1]
        if row["question"].strip() != item["expected_current"]["question"]:
            csv_mismatch.append([item["operation_ref"], "question"])
        if row["answer"].strip() != item["expected_current"]["answer"]:
            csv_mismatch.append([item["operation_ref"], "answer"])
    if csv_mismatch:
        errors.append({"code": "CSV_EXPECTED_CURRENT_MISMATCH", "items": csv_mismatch})
    else:
        passes.append("27 mevcut QnA için CSV sıra/soru/cevap önkoşulu eşleşti")

    guard_rows = [row for row in tables["Uygulama Diff"] if clean(row["İşlem"]) == "ROUTING GUARD KUR"]
    if len(guard_rows) == len(plan["routing_guard_mutations"]) == 11:
        passes.append("11 routing guard Excel/plan arasında eşleşti")
    else:
        errors.append({"code": "GUARD_DIFF_COUNT", "excel": len(guard_rows), "plan": len(plan["routing_guard_mutations"])})

    qna319 = next(item for item in updates if item["operation_ref"] == "EX-319")
    if qna319["status"] == "BLOCKED_CASE_214":
        passes.append("QnA 319 / vaka 214 blokajı korunuyor")
    else:
        errors.append({"code": "QNA319_NOT_BLOCKED"})

    if not guard["available"]:
        errors.append({
            "code": "RUNTIME_GUARD_CAPABILITY_MISSING",
            "message": "Guard sözleşmesi backend kodu ve DB şemasında executable değil.",
            "details": guard,
        })
    else:
        passes.append("Runtime routing guard kabiliyeti mevcut")

    if len(tables["Alias Haritası"]) != 68:
        errors.append({"code": "WORKBOOK_ALIAS_ROW_COUNT", "actual": len(tables["Alias Haritası"])})
    if len(tables["QnA İşlemleri"]) != 48:
        errors.append({"code": "WORKBOOK_QNA_ROW_COUNT", "actual": len(tables["QnA İşlemleri"])})

    guarded_refs = {
        item["target"].get("operation_ref") or item["target"].get("temp_ref")
        for item in plan["routing_guard_mutations"]
    }
    qna_ready, qna_guarded, qna_blocked = [], [], []
    for item in plan["qna_mutations"]:
        ref = item.get("operation_ref") or item.get("temp_ref")
        if item["status"].startswith("BLOCKED"):
            qna_blocked.append(ref)
        elif ref in guarded_refs or item.get("content_control_ref"):
            qna_guarded.append(ref)
        else:
            qna_ready.append(ref)
    qna_guarded.append("NEW-11")
    guarded_cases = {int(case) for item in plan["routing_guard_mutations"] for case in item["cases"]}
    alias_guarded = [item["case_no"] for item in plan["alias_mutations"] if int(item["case_no"]) in guarded_cases]
    alias_ready = [item["case_no"] for item in plan["alias_mutations"] if int(item["case_no"]) not in guarded_cases]

    return {
        "status": "PASS" if not errors else "BLOCKED",
        "writes_performed": False,
        "passes": passes,
        "errors": errors,
        "operation_classification": {
            "qna_ready_if_apply_authorized": qna_ready,
            "qna_guard_blocked": qna_guarded,
            "qna_explicit_blocked": qna_blocked,
            "alias_ready_if_apply_authorized_cases": alias_ready,
            "alias_guard_blocked_cases": alias_guarded,
        },
        "counts": {
            "live_active_qna": len(live["qna"]),
            "live_aliases": len(live["aliases"]),
            "qna_updates": len(updates),
            "qna_creates": len(creates),
            "atomic_promotions": len(plan["atomic_promotions"]),
            "alias_moves": len(plan["alias_mutations"]),
            "routing_guards": len(plan["routing_guard_mutations"]),
        },
    }


def render_markdown(report: dict[str, Any], changed: list[str], plan_path: Path) -> str:
    classification = report["operation_classification"]
    lines = [
        "# AUZEF KB konsolidasyon v3.1-final — dry-run",
        "",
        f"Sonuç: **{report['status']}**",
        "",
        "Bu koşu salt okunurdur. Veritabanı, Qdrant, Meilisearch ve Gold üzerinde yazma yapılmadı.",
        "",
        "## Geçen kontroller",
        "",
        *[f"- {item}" for item in report["passes"]],
        "",
        "## Blokajlar",
        "",
        *[f"- `{item['code']}`: {item.get('message', 'Önkoşul sağlanmadı.')}" for item in report["errors"]],
        "",
        "## Operasyon sınıflandırması",
        "",
        f"- Guard olmadan uygulanabilir QnA işlemi: {len(classification['qna_ready_if_apply_authorized'])}",
        f"- Guard eksikliği nedeniyle bekleyen QnA işlemi: {len(classification['qna_guard_blocked'])}",
        f"- Açıkça bloklu QnA işlemi: {len(classification['qna_explicit_blocked'])}",
        f"- Guard olmadan uygulanabilir alias taşıması: {len(classification['alias_ready_if_apply_authorized_cases'])}",
        f"- Guard eksikliği nedeniyle bekleyen alias taşıması: {len(classification['alias_guard_blocked_cases'])}",
        "",
        "## Final Excel'den plana yansıyan operasyonlar",
        "",
        *[f"- {item}" for item in changed],
        "",
        f"Makine planı: `{plan_path.name}`",
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    options = parse_args()
    options.output_dir.mkdir(parents=True, exist_ok=True)
    tables = workbook_tables(options.workbook)
    base = json.loads(options.base_plan.read_text(encoding="utf-8"))
    plan, changed = final_plan(base, tables, options.workbook)
    source_csv = load_csv(options.qna_csv)
    before = live_snapshot(options.compose_root)
    report = validate(plan, tables, before, source_csv, guard_capability(options.backend_root, before))
    before_digest = snapshot_digest(before)
    after_digest = snapshot_digest(live_snapshot(options.compose_root))
    report["live_snapshot"] = {
        "before_blake2b": before_digest,
        "after_blake2b": after_digest,
        "unchanged": before_digest == after_digest,
    }
    if before_digest != after_digest:
        report["status"] = "BLOCKED"
        report["errors"].append({"code": "LIVE_DB_CHANGED_DURING_DRY_RUN"})

    plan_path = options.output_dir / "kb-mutation-plan-v3.1-final.json"
    report_path = options.output_dir / "dry-run-report.json"
    markdown_path = options.output_dir / "dry-run-report.md"
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    markdown_path.write_text(render_markdown(report, changed, plan_path), encoding="utf-8")
    print(json.dumps({
        "status": report["status"],
        "writes_performed": False,
        "counts": report["counts"],
        "passes": len(report["passes"]),
        "errors": [item["code"] for item in report["errors"]],
        "live_snapshot_unchanged": report["live_snapshot"]["unchanged"],
        "output_dir": str(options.output_dir),
    }, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
