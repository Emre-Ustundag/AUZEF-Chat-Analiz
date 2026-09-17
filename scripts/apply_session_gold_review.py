"""İnsan inceleme kararlarını deterministik olarak reviewed Gold katmanına uygular.

Girdi: dondurulmuş Gold v2 (değiştirilmez) + doldurulmuş review Excel'i.
Çıktı: ``gold-v2-reviewed-*`` katmanı ve karar denetim izi.

Kurallar:
- Yalnız izin verilen enum kararlar kabul edilir.
- Her karar mevcut Gold durumuna karşı önkoşul kontrolünden geçer.
- Karar dışında hiçbir alan değişmez; eski frozen Gold dosyaları yazılmaz.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from scripts.build_session_gold_review_package import (  # noqa: E402
    CONTEXT_DECISIONS,
    INTENT_DECISIONS,
    load_split_texts,
)

GOLD_DIR = ROOT / "outputs" / "gold-v2-final-20260917"
BASELINE = ROOT / "outputs" / "kb-migration-v3.1-local-apply-20260917" / "baseline"
LAYER_VERSION = "gold-v2-reviewed"
STATUS_EXCLUDED = "EXCLUDED_FROM_EVAL"
STATUS_HOLD = "SOURCE_MISSING_HOLD"
STATUS_PENDING = "PENDING_CONTENT"
OUTPUT_FILES = (
    "gold-v2-reviewed-ready.jsonl",
    "gold-v2-reviewed-excluded.jsonl",
    "gold-v2-reviewed-hold.jsonl",
    "gold-v2-reviewed-all.jsonl",
    "review-decisions.jsonl",
)


class ReviewError(RuntimeError):
    """Review dosyası sözleşmeye uymuyor; hiçbir karar uygulanmaz."""


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def jsonl(rows: list[dict[str, Any]]) -> str:
    return "".join(json.dumps(r, ensure_ascii=False, sort_keys=True) + "\n" for r in rows)


def text(value: Any) -> str:
    return str(value if value is not None else "").strip()


def parse_ids(value: Any) -> list[int]:
    raw = text(value)
    if not raw:
        return []
    ids = []
    for token in raw.replace(";", ",").split(","):
        token = token.strip()
        if not token:
            continue
        if not token.isdigit():
            raise ReviewError(f"Geçersiz QnA ID: {token!r}")
        ids.append(int(token))
    return ids


def read_review(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True)
    try:
        decisions: dict[int, dict[str, Any]] = {}

        def rows(sheet_name: str) -> list[dict[str, Any]]:
            sheet = workbook[sheet_name]
            values = list(sheet.iter_rows(values_only=True))
            header = [text(cell) for cell in values[0]]
            return [dict(zip(header, row, strict=False)) for row in values[1:] if row and row[0] is not None]

        for row in rows("1-Context Review"):
            case_id = int(row["Vaka"])
            decision = text(row["Context Decision"])
            if decision not in CONTEXT_DECISIONS:
                raise ReviewError(f"Vaka {case_id}: izin verilmeyen context kararı {decision!r}")
            decisions[case_id] = {
                "case_id": case_id, "context_decision": decision, "intent_decision": None,
                "final_expected_qna_ids": parse_ids(row.get("Final expected QnA IDs")),
                "reviewer_note": text(row.get("Reviewer note")), "sheet": "1-Context Review",
            }
        for row in rows("2-Intent Review"):
            case_id = int(row["Vaka"])
            decision = text(row["Intent Decision"])
            if decision not in INTENT_DECISIONS:
                raise ReviewError(f"Vaka {case_id}: izin verilmeyen intent kararı {decision!r}")
            entry = decisions.setdefault(case_id, {"case_id": case_id, "context_decision": None,
                                                   "intent_decision": None, "final_expected_qna_ids": [],
                                                   "reviewer_note": "", "sheet": "2-Intent Review"})
            entry.update(intent_decision=decision,
                         final_expected_qna_ids=parse_ids(row.get("Final expected QnA IDs")),
                         reviewer_note=text(row.get("Reviewer note")) or entry["reviewer_note"])
        anomaly = []
        for row in rows("3-Data Anomaly"):
            anomaly.append({"case_id": int(row["Vaka"]), "reviewer_note": text(row.get("Reviewer note")),
                            "context_decision": text(row.get("Context Decision")),
                            "intent_decision": text(row.get("Intent Decision"))})
        return {"decisions": decisions, "anomaly": anomaly}
    finally:
        workbook.close()


def apply_decisions(gold: dict[int, dict[str, Any]], review: dict[str, Any], split_texts: dict[int, dict[str, Any]],
                    active_qna: set[int], review_source: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    records: dict[int, dict[str, Any]] = {case_id: json.loads(json.dumps(record)) for case_id, record in gold.items()}
    audit: list[dict[str, Any]] = []

    for case_id, decision in sorted(review["decisions"].items()):
        if case_id not in records:
            raise ReviewError(f"Vaka {case_id} Gold v2'de yok")
        record = records[case_id]
        previous_status = record["status"]
        previous_ids = list(record["expected_qna_ids"])
        entry = {
            "case_id": case_id, "review_source": review_source, "sheet": decision["sheet"],
            "context_decision": decision["context_decision"], "intent_decision": decision["intent_decision"],
            "reviewer_note": decision["reviewer_note"], "previous_status": previous_status,
            "previous_expected_qna_ids": previous_ids,
        }
        unknown = [qid for qid in decision["final_expected_qna_ids"] if qid not in active_qna]
        if unknown:
            raise ReviewError(f"Vaka {case_id}: baseline'da olmayan QnA ID {unknown}")

        if decision["context_decision"]:
            if previous_status != "CONTEXT_REQUIRED":
                raise ReviewError(f"Vaka {case_id}: context kararı için beklenen durum CONTEXT_REQUIRED, bulunan {previous_status}")
            if not decision["reviewer_note"]:
                raise ReviewError(f"Vaka {case_id}: context kararı için gerekçe zorunlu")
            if decision["context_decision"] == "EXCLUDE_FROM_EVAL":
                record.update(status=STATUS_EXCLUDED, expected_qna_ids=[], expected_intents=[], expected_qna_refs=[],
                              exclusion_reason=decision["reviewer_note"])
                action = "evaluation hedefi olmaktan çıkarıldı; veri evreninde kaldı"
            elif decision["context_decision"] == "SOURCE_MISSING_HOLD":
                record.update(status=STATUS_HOLD, expected_qna_ids=[], expected_intents=[], expected_qna_refs=[],
                              hold_reason=decision["reviewer_note"], hold_type="SOURCE_MISSING")
                action = "kaynak bulunamadı; ayrı hold kümesine alındı (214'ten farklı)"
            elif decision["context_decision"] == "KEEP_CONTEXT_REQUIRED":
                action = "değişiklik yok"
            else:
                action = "yeniden sınıflandırıldı"
                record.update(status="READY")
            entry["applied_action"] = action

        if decision["intent_decision"]:
            if previous_status != "READY":
                raise ReviewError(f"Vaka {case_id}: intent kararı için beklenen durum READY, bulunan {previous_status}")
            final_ids = decision["final_expected_qna_ids"]
            if not final_ids:
                raise ReviewError(f"Vaka {case_id}: intent kararı için beklenen QnA ID gerekli")
            if decision["intent_decision"] == "KEEP_SINGLE_INTENT":
                if final_ids != previous_ids:
                    raise ReviewError(f"Vaka {case_id}: tek niyet kararında ID değişimi ({previous_ids} → {final_ids})")
                entry["applied_action"] = "insan kararı korundu (tek niyet)"
            elif decision["intent_decision"] == "MARK_MULTI_INTENT":
                pieces = split_texts.get(case_id, {}).get("pieces", [])
                if len(pieces) != len(final_ids):
                    raise ReviewError(f"Vaka {case_id}: {len(final_ids)} hedef için {len(pieces)} niyet parçası var")
                record["expected_intents"] = [
                    {"intent_index": index, "intent_text": piece, "accepted_qna_ids": [qid],
                     "accepted_answers": [{"qna_id": qid}]}
                    for index, (piece, qid) in enumerate(zip(pieces, final_ids, strict=True), start=1)
                ]
                record.update(expected_qna_ids=sorted(final_ids), multi_intent=True,
                              expected_qna_refs=[f"QNA-{qid}" for qid in sorted(final_ids)])
                entry["applied_action"] = "çoklu niyet olarak işaretlendi; her parça kendi QnA hedefine bağlandı"
            else:
                record.update(status=STATUS_EXCLUDED, expected_qna_ids=[], expected_intents=[], expected_qna_refs=[],
                              exclusion_reason=decision["reviewer_note"])
                entry["applied_action"] = "intent incelemesi sonucu test dışı"

        record["review"] = {
            "context_decision": decision["context_decision"], "intent_decision": decision["intent_decision"],
            "reviewer_note": decision["reviewer_note"], "review_source": review_source,
            "previous_status": previous_status, "previous_expected_qna_ids": previous_ids,
            "applied_action": entry.get("applied_action"),
        }
        entry.update(new_status=record["status"], new_expected_qna_ids=record["expected_qna_ids"],
                     multi_intent=record["multi_intent"])
        audit.append(entry)

    for note in review["anomaly"]:
        audit.append({
            "case_id": note["case_id"], "review_source": review_source, "sheet": "3-Data Anomaly",
            "context_decision": note["context_decision"] or "N/A", "intent_decision": note["intent_decision"] or "N/A",
            "reviewer_note": note["reviewer_note"],
            "previous_status": gold[note["case_id"]]["status"], "new_status": gold[note["case_id"]]["status"],
            "previous_expected_qna_ids": gold[note["case_id"]]["expected_qna_ids"],
            "new_expected_qna_ids": gold[note["case_id"]]["expected_qna_ids"],
            "multi_intent": gold[note["case_id"]]["multi_intent"],
            "applied_action": "annotation değiştirilmedi; teknik normalizasyon düzeltmesiyle çözülüyor",
        })
    return [records[k] for k in sorted(records)], audit


def build(review_path: Path, output_dir: Path, created_at: str) -> dict[str, Any]:
    gold = {int(json.loads(line)["case_id"]): json.loads(line)
            for line in (GOLD_DIR / "gold-v2-all.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()}
    active_qna = {int(row["id"]) for row in json.loads((BASELINE / "qna-canonical.json").read_text(encoding="utf-8"))
                  if row["status"] == 1}
    review = read_review(review_path)
    # Kaynak referansı çağrı biçiminden bağımsız olsun: repo içindeyse göreli yol.
    resolved = review_path.resolve()
    source = str(resolved.relative_to(ROOT)) if resolved.is_relative_to(ROOT) else str(resolved)
    records, audit = apply_decisions(gold, review, load_split_texts(), active_qna, source)

    by_status: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        by_status.setdefault(record["status"], []).append(record)
    counts = {status: len(items) for status, items in sorted(by_status.items())}
    manifest = {
        "layer": LAYER_VERSION,
        "created_at": created_at,
        "source_gold_manifest_sha256": sha256(GOLD_DIR / "gold-v2-manifest.json"),
        "source_gold_all_sha256": sha256(GOLD_DIR / "gold-v2-all.jsonl"),
        "review_workbook": {"name": review_path.name, "sha256": sha256(review_path)},
        "kb_canonical_baseline_sha256": sha256(BASELINE / "qna-canonical.json"),
        "counts": counts,
        "decisions": {
            "context": {entry["case_id"]: entry["context_decision"] for entry in audit if entry.get("context_decision") not in (None, "N/A")},
            "intent": {entry["case_id"]: entry["intent_decision"] for entry in audit if entry.get("intent_decision") not in (None, "N/A")},
        },
        "multi_intent_cases": sorted(r["case_id"] for r in records if r["multi_intent"]),
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "gold-v2-reviewed-ready.jsonl": jsonl(by_status.get("READY", []) + by_status.get("CONTEXT_REQUIRED", [])),
        "gold-v2-reviewed-excluded.jsonl": jsonl(by_status.get(STATUS_EXCLUDED, [])),
        "gold-v2-reviewed-hold.jsonl": jsonl(by_status.get(STATUS_HOLD, []) + by_status.get(STATUS_PENDING, [])),
        "gold-v2-reviewed-all.jsonl": jsonl(records),
        "review-decisions.jsonl": jsonl(audit),
    }
    for name, content in files.items():
        (output_dir / name).write_text(content, encoding="utf-8")
    manifest["outputs_sha256"] = {name: sha256(output_dir / name) for name in OUTPUT_FILES}
    (output_dir / "gold-v2-reviewed-manifest.json").write_text(dump(manifest), encoding="utf-8")
    return {"manifest": manifest, "records": records, "audit": audit}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--review-workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs" / "gold-v2-reviewed-final-20260917")
    parser.add_argument("--created-at", required=True)
    return parser.parse_args()


def main() -> int:
    options = parse_args()
    result = build(options.review_workbook, options.output_dir, options.created_at)
    print(json.dumps({"counts": result["manifest"]["counts"],
                      "multi_intent_cases": result["manifest"]["multi_intent_cases"],
                      "decisions": result["manifest"]["decisions"],
                      "output_dir": str(options.output_dir)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
