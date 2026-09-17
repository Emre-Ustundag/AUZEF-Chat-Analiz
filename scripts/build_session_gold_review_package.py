"""Session-gold v2 için kompakt insan inceleme paketi üretir (salt okunur).

Kapsam: çözülemeyen 8 bağlam vakası, 7 multi-intent adayı ve 1 veri anomalisi.
Hiçbir gold/status etiketi değiştirilmez; yalnız kanıt toplanır ve karar
alanları hazırlanır. Kararlar ayrı bir görevde uygulanacak.

Kanıt kaynakları: frozen Gold v2, session-gold v2 çıktıları, ham oturum
çıkarımı (PII maskeli) ve insan inceleme workbook'unun split sayfası.
"""

from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from scripts.build_session_gold_v2 import (  # noqa: E402
    EXTRACT_DIR,
    GOLD_DIR,
    PENDING_CASE,
    SOURCES,
    choose_occurrence,
    dedup_messages,
    gap_statistics,
    is_conversational,
    locate_target,
    normalize,
    parse_time,
    read_jsonl,
    segment_session,
)

SESSION_GOLD_DIR = ROOT / "outputs" / "session-gold-v2-final-20260917"
REVIEW_WORKBOOK = SOURCES / "yanit-gold-inceleme-174.xlsx"
CONTEXT_CASES = (106, 199, 237, 248, 319, 416, 457, 514)
INTENT_CASES = (62, 71, 212, 318, 436, 456, 480)
FROZEN_FILES = (
    GOLD_DIR / "gold-v2-ready.jsonl",
    GOLD_DIR / "gold-v2-context-required.jsonl",
    GOLD_DIR / "gold-v2-pending.jsonl",
    GOLD_DIR / "gold-v2-all.jsonl",
    GOLD_DIR / "gold-v2-manifest.json",
    SESSION_GOLD_DIR / "session-gold-v2.jsonl",
    SESSION_GOLD_DIR / "session-targets.jsonl",
    SESSION_GOLD_DIR / "session-gold-manifest.json",
)
CONTEXT_DECISIONS = ("KEEP_CONTEXT_REQUIRED", "RECLASSIFY_READY_STANDALONE", "EXCLUDE_FROM_EVAL", "SOURCE_MISSING_HOLD")
INTENT_DECISIONS = ("KEEP_SINGLE_INTENT", "MARK_MULTI_INTENT", "EXCLUDE_FROM_EVAL")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="E2EFDA")
HEADER = PatternFill("solid", fgColor="D9D9D9")
EVIDENCE_AFTER_TARGET = 2


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def turn_view(message: dict[str, Any], *, evidence_only: bool = False) -> dict[str, Any]:
    return {
        "message_id": message["message_id"],
        "role": "user" if message["direction"] == "Kullanıcı" else "assistant",
        "timestamp": message["time"],
        "text": message["text"],
        "conversational": is_conversational(message),
        "message_type": message["message_type"],
        "evidence_only_not_evaluator_context": evidence_only,
    }


def minutes_between(first: str, second: str) -> float | None:
    start, end = parse_time(first), parse_time(second)
    return round((end - start).total_seconds() / 60.0, 2) if start and end else None


def fuzzy_candidates(message: str, corpus: list[tuple[str, str]], limit: int = 3) -> list[dict[str, Any]]:
    """Kaba benzerlik; yalnız insan incelemesine yardımcı kanıttır."""
    wanted = normalize(message)
    tokens = {token for token in wanted.split() if len(token) > 3}
    scored = []
    for key, text in corpus:
        candidate = normalize(text)
        if tokens and not tokens & set(candidate.split()):
            continue
        scored.append((round(difflib.SequenceMatcher(None, wanted, candidate).ratio(), 3), key, text))
    scored.sort(key=lambda item: (-item[0], item[1]))
    return [{"score": score, "ref": key, "text": text[:400]} for score, key, text in scored[:limit]]


def collect(session_corpus: Path | None) -> dict[str, Any]:
    gold = {int(r["case_id"]): r for r in read_jsonl(GOLD_DIR / "gold-v2-all.jsonl")}
    targets = {int(r["case_id"]): r for r in read_jsonl(SESSION_GOLD_DIR / "session-targets.jsonl")}
    context_rows = {int(r["case_id"]): r for r in read_jsonl(SESSION_GOLD_DIR / "context-resolutions.jsonl")}
    report = json.loads((SESSION_GOLD_DIR / "session-gold-report.json").read_text(encoding="utf-8"))
    alias = {r["alias_id"]: r for r in read_jsonl(SOURCES / "alias-session-matches.jsonl")}
    sessions = read_jsonl(EXTRACT_DIR / "sessions-extract.jsonl")
    raw_by_session = {s["session_id"]: s for s in sessions}
    messages_by_session, _ = dedup_messages(sessions)
    threshold = gap_statistics(messages_by_session)["selected_threshold_minutes"]
    conversational = {sid: [m for m in msgs if is_conversational(m)] for sid, msgs in messages_by_session.items()}
    segments_by_session = {sid: segment_session(msgs, threshold) for sid, msgs in conversational.items()}

    anomaly_cases = tuple(report["audit"]["gold_message_fragment_cases"])
    baseline = json.loads((ROOT / "outputs" / "kb-migration-v3.1-local-apply-20260917" / "baseline" / "qna-canonical.json").read_text(encoding="utf-8"))
    aliases_baseline = json.loads((ROOT / "outputs" / "kb-migration-v3.1-local-apply-20260917" / "baseline" / "qna-aliases.json").read_text(encoding="utf-8"))
    kb_corpus = [(f"QNA-{r['id']}", r["question_text"]) for r in baseline if r["status"] == 1]
    kb_corpus += [(f"QNA-{r['qna_id']}", r["query_text"]) for r in aliases_baseline]

    split_texts = load_split_texts()
    items: dict[int, dict[str, Any]] = {}

    def item(case_id: int, review_type: str) -> dict[str, Any]:
        record = items.setdefault(case_id, {
            "case_id": case_id,
            "review_types": [],
            "gold_user_message": gold[case_id]["user_message"],
            "current_status": gold[case_id]["status"],
            "current_expected_qna_ids": gold[case_id]["expected_qna_ids"],
            "current_expected_qna_refs": gold[case_id]["expected_qna_refs"],
            "gold_notes": gold[case_id]["notes"],
            "source_refs": {
                "gold": str((GOLD_DIR / "gold-v2-all.jsonl").relative_to(ROOT)),
                "sessions_extract": str((EXTRACT_DIR / "sessions-extract.jsonl").relative_to(ROOT)),
                "session_gold_report": str((SESSION_GOLD_DIR / "session-gold-report.json").relative_to(ROOT)),
            },
            "allowed_decisions": {"context": list(CONTEXT_DECISIONS), "intent": list(INTENT_DECISIONS)},
        })
        if review_type not in record["review_types"]:
            record["review_types"].append(review_type)
        return record

    # --- A: çözülemeyen bağlam vakaları ---------------------------------
    for case_id in CONTEXT_CASES:
        record = item(case_id, "UNRESOLVED_CONTEXT")
        context_row = context_rows[case_id]
        match = alias[gold[case_id]["alias_id"]]
        occurrence, policy = choose_occurrence(match.get("occurrences") or [], prefer_context=True)
        evidence: dict[str, Any] = {
            "review_reason": context_row["resolution_status"],
            "why_context_required": context_row["why_context_required"],
            "review_note_source": context_row["source_evidence"].get("review_session_source"),
            "match_status": match.get("match_status"),
            "occurrence_count": len(match.get("occurrences") or []),
            "assignment_policy": policy,
            "session_id": occurrence["session_id"] if occurrence else None,
            "segmentation_threshold_minutes": threshold,
        }
        if occurrence:
            session_id = occurrence["session_id"]
            segments = segments_by_session.get(session_id, [])
            located = locate_target(segments, int(occurrence["turn_index"]), gold[case_id]["user_message"])
            raw_messages = raw_by_session[session_id]["messages"]
            if located:
                segment_index, position, method = located
                segment = segments[segment_index]
                target = segment[position]
                evidence.update({
                    "match_method": method,
                    "segment_index": segment_index,
                    "segment_count_in_session": len(segments),
                    "target_turn_id": target["message_id"],
                    "target_timestamp": target["time"],
                    "target_text": target["text"],
                    "turns_before_target_in_segment": [turn_view(m) for m in segment[:position]],
                    "turns_after_target_evidence_only": [turn_view(m, evidence_only=True)
                                                         for m in segment[position + 1:position + 1 + EVIDENCE_AFTER_TARGET]],
                    "previous_segment_last_turns": [turn_view(m, evidence_only=True)
                                                    for m in (segments[segment_index - 1][-3:] if segment_index else [])],
                    "previous_gap_minutes": (minutes_between(segments[segment_index - 1][-1]["time"], target["time"])
                                             if segment_index else None),
                    "next_gap_minutes": (minutes_between(target["time"], segment[position + 1]["time"])
                                         if position + 1 < len(segment) else None),
                })
                raw_before = [m for m in raw_messages
                              if (m["message_order"] or 0) < (target["message_order"] or 0)]
                evidence.update({
                    "raw_messages_before_target_total": len(raw_before),
                    "raw_non_conversational_before_target": [turn_view(m, evidence_only=True) for m in raw_before
                                                             if not is_conversational(m)][-5:],
                    "filtered_user_messages_before_target": [turn_view(m, evidence_only=True) for m in raw_before
                                                             if m["direction"] == "Kullanıcı" and not m["is_gold_user_turn"]][-5:],
                    "conversational_before_target_in_session": sum(
                        1 for m in raw_before if is_conversational(m)),
                })
        else:
            corpus = []
            if session_corpus and session_corpus.exists():
                for line in session_corpus.open(encoding="utf-8"):
                    row = json.loads(line)
                    for index, turn in enumerate(row["text"].split(" / "), start=1):
                        corpus.append((f"{row['session']}#{index}", turn))
            evidence.update({
                "normalized_gold_message": normalize(gold[case_id]["user_message"]),
                "searched_sources": [
                    str((EXTRACT_DIR / "sessions-extract.jsonl").relative_to(ROOT)),
                    str(session_corpus) if session_corpus else None,
                    str((SOURCES / "alias-session-matches.jsonl").relative_to(ROOT)),
                ],
                "exact_match_found": False,
                "fuzzy_session_candidates": fuzzy_candidates(gold[case_id]["user_message"], corpus),
                "why_not_found": "Gold mesajı hiçbir oturumun kullanıcı turn'üyle birebir eşleşmiyor "
                                 "(alias eşleşmesi de 'not_found'); mesaj kırpılmış görünüyor.",
            })
        record["evidence_context"] = evidence

    # --- B: multi-intent adayları ---------------------------------------
    for case_id in INTENT_CASES:
        record = item(case_id, "MULTI_INTENT_REVIEW")
        parts = split_texts.get(case_id, {})
        pieces = parts.get("pieces", [])
        record["evidence_intent"] = {
            "review_reason": "Splitter 'Gerekli' dedi; Gold tek hedefle donduruldu",
            "splitter_label": gold[case_id]["split_audit"].get("split_label"),
            "splitter_note": parts.get("note"),
            "human_review_decision": parts.get("human_decision") or gold[case_id]["source_decision"].get("decision"),
            "gold_split_audit": gold[case_id]["split_audit"],
            "intent_pieces": [
                {"index": index, "text": piece,
                 "kb_suggestions_heuristic_only": fuzzy_candidates(piece, kb_corpus, limit=2)}
                for index, piece in enumerate(pieces, start=1)
            ],
            "current_target_in_session": {
                "evaluation_session_id": targets[case_id]["evaluation_session_id"] if case_id in targets else None,
                "turn_type": targets[case_id]["turn_type"] if case_id in targets else None,
            },
            "second_piece_looks_like_elaboration": bool(pieces) and len(
                {tuple(sorted(gold[case_id]["expected_qna_ids"]))}) == 1 and len(pieces) > 1,
        }

    # --- C: veri anomalisi ----------------------------------------------
    for case_id in anomaly_cases:
        record = item(case_id, "DATA_ANOMALY_MERGED_MESSAGE")
        target = targets.get(case_id)
        session_id = target["source_session_id"] if target else None
        real_turn = None
        if session_id:
            segments = segments_by_session[session_id]
            segment = segments[target["segment_index"]]
            real_turn = turn_view(segment[target["turn_index"]])
        gold_message = gold[case_id]["user_message"]
        real_text = (real_turn or {}).get("text", "")
        gold_norm, real_norm = normalize(gold_message), normalize(real_text)
        turkish_dotted = "\u0307" in real_norm or "\u0307" in gold_norm
        record["evidence_anomaly"] = {
            "review_reason": ("Gold mesajı ile gerçek turn birebir eşleşmiyor; eşleşme 'içerir' kuralına düştü"),
            "verified_cause": (
                "Türkçe noktalı büyük İ casefold edilince 'i' + U+0307 (birleşen nokta) üretiyor; "
                "bu yüzden 'ÇÖZÜM MERKEZİ' ile 'Çözüm merkezi' eşit sayılmıyor. "
                "Daha önce raporlanan \" / \" birleştirme açıklaması bu vaka için geçerli DEĞİL."
                if turkish_dotted else
                "Gold mesajı gerçek turn'ün parçası; alias tabanı turn'leri ' / ' ile birleştirip yeniden bölmüş."
            ),
            "gold_message": gold_message,
            "gold_message_repr": repr(gold_message),
            "real_full_turn": real_turn,
            "real_text_repr": repr(real_text),
            "normalized_gold_codepoints": [hex(ord(ch)) for ch in gold_norm][-6:],
            "normalized_real_codepoints": [hex(ord(ch)) for ch in real_norm][-6:],
            "texts_differ_only_by_case": gold_norm.replace("\u0307", "") == real_norm.replace("\u0307", ""),
            "suggested_fix": ("Normalizasyonda Türkçe büyük harf eşlemesi (İ→i, I→ı) casefold'dan önce uygulanırsa "
                              "bu vaka birebir eşleşir; düzeltme session-gold'u yeniden üretmeyi gerektirir."),
            "match_method": target.get("target_match_method") if target else None,
            "evaluation_session_id": target["evaluation_session_id"] if target else None,
            "context_resolution_status": context_rows.get(case_id, {}).get("resolution_status"),
        }

    for record in items.values():
        record["review_types"].sort()
        record["needs_context_decision"] = "UNRESOLVED_CONTEXT" in record["review_types"]
        record["needs_intent_decision"] = "MULTI_INTENT_REVIEW" in record["review_types"]
        record["prefilled_final_expected_qna_ids"] = record["current_expected_qna_ids"]
    return {"items": [items[k] for k in sorted(items)], "threshold_minutes": threshold,
            "anomaly_cases": list(anomaly_cases)}


def load_split_texts() -> dict[int, dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(REVIEW_WORKBOOK, read_only=True, data_only=True)
    try:
        rows = list(workbook["Split vakaları"].iter_rows(min_row=4, values_only=True))
        header = [str(value or "").strip() for value in rows[0]]
        index = {name: position for position, name in enumerate(header)}
        result = {}
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            case_id = int(row[index["Vaka"]])
            result[case_id] = {
                "pieces": [part.strip() for part in str(row[index["Luna ayrımı"]] or "").split("|") if part.strip()],
                "note": str(row[index["Açıklama"]] or "").strip(),
                "human_decision": str(row[index["Model kararı"]] or "").strip(),
                "split_decision": str(row[index["Split kararı"]] or "").strip(),
            }
        return result
    finally:
        workbook.close()


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------


def style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.fill = HEADER
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_widths(sheet, widths: list[int]) -> None:
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width


def short(turns: list[dict[str, Any]], limit: int = 3) -> str:
    return "\n".join(f"[{t['role']}] {t['text'][:220]}" for t in turns[-limit:]) or "—"


def build_workbook(package: dict[str, Any], path: Path) -> dict[str, Any]:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Başlangıç"
    items = package["items"]
    context_items = [i for i in items if i["needs_context_decision"]]
    intent_items = [i for i in items if i["needs_intent_decision"]]
    anomaly_items = [i for i in items if "DATA_ANOMALY_MERGED_MESSAGE" in i["review_types"]]

    summary.append(["Session-gold v2 — insan inceleme paketi"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    summary.append(["Toplam review item", len(items)])
    summary.append(["Çözülemeyen bağlam vakası", len(context_items)])
    summary.append(["Multi-intent adayı", len(intent_items)])
    summary.append(["Veri anomalisi", len(anomaly_items)])
    summary.append(["Birden fazla review türü", len([i for i in items if len(i["review_types"]) > 1])])
    summary.append([])
    summary.append(["Tamamlanan karar", "=COUNTIF('1-Context Review'!C2:C100,\"<>\")+COUNTIF('2-Intent Review'!C2:C100,\"<>\")"])
    summary.append(["Bekleyen karar", f"={len(context_items) + len(intent_items)}-B9"])
    summary.append([])
    summary.append(["Nasıl doldurulur",
                    "Sarı hücreler doldurulacak. Karar listeden seçilir. Gri hücreler kaynak kanıttır, değiştirmeyin. "
                    "Hedeften sonraki turn'ler yalnız inceleme kanıtıdır; değerlendirmede bağlam olarak kullanılmaz."])
    set_widths(summary, [34, 90])
    for row in range(3, 8):
        summary.cell(row=row, column=2).font = Font(bold=True)

    # 1 - Context Review
    sheet = workbook.create_sheet("1-Context Review")
    headers = ["Vaka", "Neden incelemede", "Context Decision", "Final expected QnA IDs", "Reviewer note", "Durum",
               "Gold mesajı", "Gold durumu", "Mevcut QnA", "Oturum", "Hedef turn id", "Hedef zamanı",
               "Hedef öncesi gerçek turn sayısı", "Hedef öncesi turnler", "HEDEF", "Sonraki turnler (yalnız kanıt)",
               "Önceki segment son turnleri", "Önceki boşluk (dk)", "Sonraki boşluk (dk)", "Ham: hedef öncesi mesaj",
               "Ham: filtrelenen kullanıcı mesajları", "İnceleme notu"]
    sheet.append(headers)
    style_header(sheet)
    for row_index, entry in enumerate(context_items, start=2):
        evidence = entry["evidence_context"]
        before = evidence.get("turns_before_target_in_segment", [])
        sheet.append([
            entry["case_id"], evidence["review_reason"], None, ", ".join(map(str, entry["prefilled_final_expected_qna_ids"])) or "", None,
            f'=IF(AND(C{row_index}<>"",E{row_index}<>""),"TAMAM","EKSİK")',
            entry["gold_user_message"], entry["current_status"],
            ", ".join(map(str, entry["current_expected_qna_ids"])) or "—",
            evidence.get("session_id") or "BULUNAMADI", evidence.get("target_turn_id") or "—",
            evidence.get("target_timestamp") or "—", len(before), short(before),
            evidence.get("target_text") or entry["gold_user_message"],
            short(evidence.get("turns_after_target_evidence_only", []), limit=EVIDENCE_AFTER_TARGET),
            short(evidence.get("previous_segment_last_turns", [])),
            evidence.get("previous_gap_minutes"), evidence.get("next_gap_minutes"),
            evidence.get("raw_messages_before_target_total", 0),
            short(evidence.get("filtered_user_messages_before_target", [])),
            evidence["why_context_required"],
        ])
        for column in ("C", "D", "E"):
            sheet[f"{column}{row_index}"].fill = YELLOW
        if evidence["review_reason"] == "SOURCE_NOT_FOUND":
            sheet[f"J{row_index}"].fill = RED
        if evidence.get("turns_after_target_evidence_only"):
            sheet[f"P{row_index}"].fill = RED
        if evidence.get("previous_segment_last_turns"):
            sheet[f"Q{row_index}"].fill = GREEN
    set_widths(sheet, [7, 26, 24, 22, 28, 10, 52, 16, 14, 38, 14, 20, 12, 60, 52, 52, 52, 14, 14, 12, 40, 60])
    sheet.freeze_panes = "C2"
    validation = DataValidation(type="list", formula1='"' + ",".join(CONTEXT_DECISIONS) + '"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"C2:C{len(context_items) + 1}")

    # 2 - Intent Review
    sheet = workbook.create_sheet("2-Intent Review")
    headers = ["Vaka", "Intent Decision", "Final expected QnA IDs", "Reviewer note", "Durum", "Gold mesajı",
               "Mevcut QnA", "Splitter etiketi", "İnsan kararı", "Parça 1", "Parça 1 — KB önerisi (sezgisel)",
               "Parça 2", "Parça 2 — KB önerisi (sezgisel)", "Splitter açıklaması", "Oturumdaki turn tipi"]
    sheet.append(headers)
    style_header(sheet)
    for row_index, entry in enumerate(intent_items, start=2):
        evidence = entry["evidence_intent"]
        pieces = evidence["intent_pieces"]

        def suggestion(index: int) -> str:
            if index >= len(pieces):
                return "—"
            return "\n".join(f"{c['ref']} ({c['score']}) {c['text'][:90]}" for c in pieces[index]["kb_suggestions_heuristic_only"]) or "—"

        sheet.append([
            entry["case_id"], None, ", ".join(map(str, entry["prefilled_final_expected_qna_ids"])), None,
            f'=IF(AND(B{row_index}<>"",D{row_index}<>""),"TAMAM","EKSİK")',
            entry["gold_user_message"], ", ".join(map(str, entry["current_expected_qna_ids"])),
            evidence["splitter_label"], evidence["human_review_decision"],
            pieces[0]["text"] if pieces else "—", suggestion(0),
            pieces[1]["text"] if len(pieces) > 1 else "—", suggestion(1),
            evidence["splitter_note"], evidence["current_target_in_session"]["turn_type"],
        ])
        for column in ("B", "C", "D"):
            sheet[f"{column}{row_index}"].fill = YELLOW
    set_widths(sheet, [7, 24, 22, 28, 10, 60, 14, 16, 22, 52, 46, 52, 46, 34, 34])
    sheet.freeze_panes = "B2"
    validation = DataValidation(type="list", formula1='"' + ",".join(INTENT_DECISIONS) + '"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"B2:B{len(intent_items) + 1}")

    # 3 - Data Anomaly
    sheet = workbook.create_sheet("3-Data Anomaly")
    sheet.append(["Vaka", "Context Decision", "Intent Decision", "Final expected QnA IDs", "Reviewer note", "Durum",
                  "Anomali", "Doğrulanan neden", "Gold mesajı", "Gerçek tam mesaj", "Yalnız büyük/küçük harf farkı",
                  "Önerilen düzeltme", "Eşleşme yöntemi", "Evaluation session", "Bağlam çözümü"])
    style_header(sheet)
    for row_index, entry in enumerate(anomaly_items, start=2):
        evidence = entry["evidence_anomaly"]
        real = evidence.get("real_full_turn") or {}
        sheet.append([
            entry["case_id"],
            "N/A" if not entry["needs_context_decision"] else None,
            "N/A" if not entry["needs_intent_decision"] else None,
            ", ".join(map(str, entry["prefilled_final_expected_qna_ids"])) or "", None,
            f'=IF(E{row_index}<>"","TAMAM","EKSİK")',
            evidence["review_reason"], evidence["verified_cause"], evidence["gold_message"], real.get("text", "—"),
            "EVET" if evidence["texts_differ_only_by_case"] else "hayır", evidence["suggested_fix"],
            evidence["match_method"], evidence["evaluation_session_id"], evidence["context_resolution_status"],
        ])
        for column in ("B", "C", "D", "E"):
            sheet[f"{column}{row_index}"].fill = YELLOW
        sheet[f"I{row_index}"].fill = RED
        sheet[f"J{row_index}"].fill = RED
    set_widths(sheet, [7, 22, 22, 22, 28, 10, 48, 70, 26, 34, 18, 60, 34, 40, 26])

    # Kaynak Kanıt
    sheet = workbook.create_sheet("Kaynak Kanıt")
    sheet.append(["Vaka", "Tür", "Sıra", "Rol", "Zaman", "Mesaj", "Yalnız kanıt (evaluator bağlamı değil)"])
    style_header(sheet)
    for entry in items:
        evidence = entry.get("evidence_context") or {}
        blocks = [("hedef öncesi", evidence.get("turns_before_target_in_segment", [])),
                  ("hedef sonrası (kanıt)", evidence.get("turns_after_target_evidence_only", [])),
                  ("önceki segment", evidence.get("previous_segment_last_turns", [])),
                  ("ham: filtrelenen kullanıcı", evidence.get("filtered_user_messages_before_target", [])),
                  ("ham: arayüz olayı", evidence.get("raw_non_conversational_before_target", []))]
        for label, turns in blocks:
            for position, turn in enumerate(turns, start=1):
                sheet.append([entry["case_id"], label, position, turn["role"], turn["timestamp"], turn["text"][:500],
                              "EVET" if turn["evidence_only_not_evaluator_context"] else "hayır"])
        for candidate in evidence.get("fuzzy_session_candidates", []):
            sheet.append([entry["case_id"], "fuzzy aday (kanıt)", candidate["score"], "—", "—", candidate["text"][:500], "EVET"])
    set_widths(sheet, [7, 28, 8, 12, 20, 110, 34])

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return {"context_items": len(context_items), "intent_items": len(intent_items), "anomaly_items": len(anomaly_items)}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs" / "session-gold-v2-review-20260917")
    parser.add_argument("--session-corpus", type=Path,
                        default=Path.home() / "Masaüstü" / "AUZEF-Chat-Analiz" / "outputs" / "analiz" / "sessions-yil.jsonl",
                        help="Kaynağı bulunamayan vaka için fuzzy arama yapılacak tam oturum korpusu (salt okunur)")
    parser.add_argument("--created-at", required=True)
    return parser.parse_args()


def main() -> int:
    options = parse_args()
    before = {str(path.relative_to(ROOT)): sha256(path) for path in FROZEN_FILES}
    package = collect(options.session_corpus)
    options.output_dir.mkdir(parents=True, exist_ok=True)
    excel_path = options.output_dir / "session-gold-human-review.xlsx"
    counts = build_workbook(package, excel_path)
    after = {str(path.relative_to(ROOT)): sha256(path) for path in FROZEN_FILES}
    audit = {
        "created_at": options.created_at,
        "review_universe": {
            "total_items": len(package["items"]),
            "context_cases": sorted(i["case_id"] for i in package["items"] if i["needs_context_decision"]),
            "intent_cases": sorted(i["case_id"] for i in package["items"] if i["needs_intent_decision"]),
            "anomaly_cases": package["anomaly_cases"],
            "multi_type_cases": sorted(i["case_id"] for i in package["items"] if len(i["review_types"]) > 1),
        },
        "frozen_inputs_sha256_before": before,
        "frozen_inputs_sha256_after": after,
        "frozen_inputs_unchanged": before == after,
        "segmentation_threshold_minutes": package["threshold_minutes"],
        "excel": {"path": str(excel_path.relative_to(ROOT)), "sheets": ["Başlangıç", "1-Context Review",
                                                                        "2-Intent Review", "3-Data Anomaly", "Kaynak Kanıt"],
                  **counts},
        "items": package["items"],
    }
    (options.output_dir / "session-gold-human-review.json").write_text(dump(audit), encoding="utf-8")
    print(json.dumps({
        "items": len(package["items"]), **counts,
        "context_cases": audit["review_universe"]["context_cases"],
        "intent_cases": audit["review_universe"]["intent_cases"],
        "anomaly_cases": audit["review_universe"]["anomaly_cases"],
        "multi_type_cases": audit["review_universe"]["multi_type_cases"],
        "frozen_inputs_unchanged": audit["frozen_inputs_unchanged"],
        "excel": str(excel_path), "json": str(options.output_dir / "session-gold-human-review.json"),
    }, ensure_ascii=False, indent=2))
    return 0 if audit["frozen_inputs_unchanged"] and PENDING_CASE not in {i["case_id"] for i in package["items"]} else 2


if __name__ == "__main__":
    raise SystemExit(main())
