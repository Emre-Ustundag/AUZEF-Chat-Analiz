"""4o-mini ve Luna-high için ortak, eşli hata denetimi (yalnız mevcut izlerden).

Benchmark tekrar koşulmaz; sistem, gold ve KB değiştirilmez. Kök nedenler
yalnız izden nesnel olarak türetilebiliyorsa atanır; aksi halde
NEEDS_HUMAN_REVIEW. Tam mesaj için pipeline'ın her zaman yaptığı
"spekülatif" seçim, bölücüden bağımsız bir karşı-olgu olarak kullanılır.
"""

from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from scripts.build_session_gold_v2 import normalize  # noqa: E402
from scripts.run_e2e_baseline import build_targets, evaluate  # noqa: E402

BENCH = ROOT / "outputs" / "e2e-baseline-session-gold-v2-20260917"
BASELINE = ROOT / "outputs" / "kb-migration-v3.1-local-apply-20260917" / "baseline"
MODELS = ("4o-mini", "luna-high")
NEAR_DUPLICATE_RATIO = 0.6
DATE_SIGNAL = re.compile(r"ne zaman|tarih|takvim|başla|son gün|bitiş|kaçında|hangi gün|dönem|vize|final|bütünleme|sınav(lar)?ı?n? ne", re.I)
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="D9D9D9")
REVIEW_DECISIONS = ("AGREE_AUTO_CAUSE", "SELECTOR_ERROR", "GOLD_WRONG", "GOLD_AMBIGUOUS", "KB_OVERLAP",
                    "KB_GAP", "SPLITTER_ERROR", "CALENDAR", "OTHER")


def sha256(path: Path) -> str:
    return hashlib.new("sha256", path.read_bytes()).hexdigest()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def jsonl(rows: list[dict[str, Any]]) -> str:
    return "".join(json.dumps(r, ensure_ascii=False, sort_keys=True) + "\n" for r in rows)


def similarity(first: str, second: str) -> float:
    return round(difflib.SequenceMatcher(None, normalize(first or ""), normalize(second or "")).ratio(), 3)


# --------------------------------------------------------------------------
# Model başına iz görünümü
# --------------------------------------------------------------------------


def trace_view(result: dict[str, Any], evaluation: dict[str, Any], accepted: set[int]) -> dict[str, Any]:
    events = result["events"]
    splits = [e for e in events if e["kind"] == "split_result"]
    sub_questions = splits[-1]["sub_questions"] if splits else []
    asks = [e for e in events if e["kind"] == "ask"]
    speculative = next((a for a in asks if a["thread"] == "worker"), asks[0] if asks else None)
    used = [speculative] if len(sub_questions) <= 1 else [a for a in asks if a["thread"] == "main"]
    used = [a for a in used if a]

    def gold_rank(ask: dict[str, Any] | None) -> dict[str, Any]:
        if not ask:
            return {"qna_rank": None, "prompt_index": None}
        qna_only = [c["qna_id"] for c in ask["candidates"] if c["qna_id"] is not None]
        qna_rank = next((i + 1 for i, q in enumerate(qna_only) if q in accepted), None)
        prompt_index = next((i + 1 for i, c in enumerate(ask["candidates"]) if c["qna_id"] in accepted), None)
        selected_rank = next((i + 1 for i, q in enumerate(qna_only) if q == ask["selected_qna_id"]), None)
        return {"qna_rank": qna_rank, "prompt_index": prompt_index, "selected_qna_rank": selected_rank,
                "calendar_candidates": sum(1 for c in ask["candidates"] if c["qna_id"] is None),
                "qna_candidates": len(qna_only)}

    return {
        "split_count": max(len(sub_questions), 1),
        "sub_questions": sub_questions if len(sub_questions) > 1 else [],
        "speculative_selected": speculative["selected_qna_id"] if speculative else None,
        "speculative_declined": bool(speculative and speculative["declined"]),
        "speculative_calendar": bool(speculative and speculative["selected_calendar"]),
        "speculative_correct": bool(speculative and speculative["selected_qna_id"] in accepted),
        "speculative_candidates": [c["qna_id"] for c in speculative["candidates"]] if speculative else [],
        "speculative_gold_rank": gold_rank(speculative),
        "used_asks": [{"question": a["question"][:300], "selected_qna_id": a["selected_qna_id"],
                       "selected_question": a.get("selected_question"), "declined": a["declined"],
                       "selected_calendar": a["selected_calendar"], "gold_rank": gold_rank(a),
                       "candidate_ids": [c["qna_id"] for c in a["candidates"]]} for a in used],
        "selected_qna_ids": evaluation["selected_qna_ids"],
        "declined_any": any(a["declined"] for a in used),
        "calendar_selected": evaluation["selected_calendar"],
        "fallback": evaluation["fallback_used"],
        "source": result["source"],
        "final_answer": (result["answer"] or "")[:1500],
        "final_answer_chars": len(result["answer"] or ""),
        "relaxed_correct": evaluation["correct"],
        "exact_correct": evaluation["exact"],
        "coarse_cause": evaluation["primary_cause"],
        "raw_select_outputs": [e.get("raw_output") for e in events if e["kind"] == "llm_select"],
    }


def fine_cause(case: dict[str, Any], view: dict[str, Any], qna: dict[int, dict[str, Any]],
               alias_owners: dict[str, set[int]], other_view: dict[str, Any]) -> tuple[str | None, list[str]]:
    """Birebir metrikte hatalı bir modelin ince kök nedeni (yalnız izden)."""
    if view["exact_correct"]:
        return None, []
    accepted = set(case["expected_qna_ids"])
    groups = case["expected_intent_groups"]
    secondary: list[str] = []
    if case["turn_type"] == "FOLLOW_UP_CONTEXT_AVAILABLE_BUT_NOT_REQUIRED":
        secondary.append("CONTEXT_AVAILABLE_BUT_UNUSED")
    if len(groups) == 1 and view["split_count"] >= 3:
        secondary.append("OVER_FRAGMENTATION")
    selections = [a["selected_qna_id"] for a in view["used_asks"] if a["selected_qna_id"] is not None]
    if len(selections) != len(set(selections)):
        secondary.append("FRAGMENT_DUPLICATION")
    if any(r is not None and not any(ch.isdigit() for ch in r) for r in view["raw_select_outputs"]):
        return "FORMAT_OR_PARSE_ERROR", secondary

    # Gevşek doğru ama birebir yanlış: fazladan cevap
    if view["relaxed_correct"]:
        if view["calendar_selected"]:
            secondary.append("CALENDAR_EXTRA_ANSWER")
        return ("FALSE_SPLIT_SINGLE_INTENT" if view["split_count"] > 1 else "NEEDS_HUMAN_REVIEW"), secondary

    if len(groups) > 1 and view["split_count"] < len(groups):
        return "MISSED_MULTI_INTENT", secondary
    if len(groups) == 1 and view["split_count"] > 1:
        if view["speculative_correct"]:
            return "FRAGMENT_LOST_CONTEXT", secondary + ["FALSE_SPLIT_SINGLE_INTENT"]
        fragment_had_gold = any(a["gold_rank"]["qna_rank"] for a in view["used_asks"])
        if not fragment_had_gold and view["speculative_gold_rank"]["qna_rank"]:
            return "FRAGMENT_LOST_CONTEXT", secondary + ["FALSE_SPLIT_SINGLE_INTENT"]
        if fragment_had_gold:
            secondary.append("FALSE_SPLIT_SINGLE_INTENT")
            return "SELECTOR_FRAGMENT_CONTEXT_LOSS", secondary
    if view["fallback"] and view["source"] == "academic_calendar":
        return "CALENDAR_ROUTING_INTERFERENCE", secondary
    gold_present = any(a["gold_rank"]["qna_rank"] for a in view["used_asks"])
    if not gold_present:
        if selections:
            secondary.append("SELECTOR_WRONG_POSITIVE")
        return "RETRIEVAL_TRUE_MISS", secondary
    if view["calendar_selected"]:
        return "SELECTOR_CALENDAR_DISTRACTION", secondary
    missed_groups = [g for g in groups if not set(g) & set(view["selected_qna_ids"])]
    if view["declined_any"] and (not selections or (len(groups) > 1 and missed_groups)):
        # Çoklu niyette bir parça için "hiçbiri" denmesi de bu sınıftır.
        if view["fallback"]:
            secondary.append("FALLBACK_ERROR")
        return "SELECTOR_NONE_WHEN_GOLD_PRESENT", secondary
    wrong = [s for s in selections if s not in accepted]
    if not wrong:
        return "NEEDS_HUMAN_REVIEW", secondary
    chosen = wrong[0]
    gold_question = " / ".join(qna[g]["question_text"] for g in sorted(accepted) if g in qna)
    chosen_question = qna.get(chosen, {}).get("question_text", "")
    if chosen in alias_owners.get(case["user_message"].strip(), set()):
        secondary.append("GOLD_REVIEW_CANDIDATE")
        return "KB_OVERLAP_AMBIGUITY", secondary
    gold_rank = min((a["gold_rank"]["qna_rank"] for a in view["used_asks"] if a["gold_rank"]["qna_rank"]), default=None)
    selected_rank = next((a["gold_rank"]["selected_qna_rank"] for a in view["used_asks"]
                          if a["selected_qna_id"] == chosen), None)
    if gold_rank and gold_rank > 5:
        secondary.append("RETRIEVAL_GOLD_LOW_RANK")
    if gold_rank and selected_rank and selected_rank < gold_rank and gold_rank >= 4:
        secondary.append("SELECTOR_POSITION_BIAS_SIGNAL")
    if other_view["exact_correct"] is False and set(other_view["selected_qna_ids"]) == set(view["selected_qna_ids"]):
        secondary.append("BOTH_MODELS_SAME_WRONG_CHOICE")
    if max(similarity(gold_question, chosen_question),
           similarity(qna.get(chosen, {}).get("answer_text", ""),
                      " ".join(qna[g]["answer_text"] for g in accepted if g in qna))) >= NEAR_DUPLICATE_RATIO:
        return "SELECTOR_NEAR_DUPLICATE_CONFUSION", secondary
    return "SELECTOR_WRONG_SEMANTIC_MATCH", secondary


def model_diff_labels(a: dict[str, Any], b: dict[str, Any]) -> list[str]:
    labels = []
    if a["split_count"] != b["split_count"]:
        labels.append("MODEL_DIFF_SPLITTER")
        wrong, right = (a, b) if not a["exact_correct"] else (b, a)
        if wrong["split_count"] > 1 and right["split_count"] == 1:
            labels.append("MODEL_DIFF_FRAGMENTATION")
    if a["declined_any"] != b["declined_any"]:
        labels.append("MODEL_DIFF_NONE_CALIBRATION")
    if a["calendar_selected"] != b["calendar_selected"]:
        labels.append("MODEL_DIFF_CALENDAR")
    same_input = (a["split_count"] == b["split_count"] == 1
                  and a["speculative_candidates"] == b["speculative_candidates"])
    if same_input and a["selected_qna_ids"] != b["selected_qna_ids"]:
        labels.append("MODEL_DIFF_SAME_INPUT_DIFFERENT_SELECTION")
        if a["exact_correct"] != b["exact_correct"]:
            labels.append("PURE_SELECTOR_MODEL_DIFFERENCE")
    if not labels:
        labels.append("MODEL_DIFF_OTHER")
    return labels


# --------------------------------------------------------------------------
# Ana üretim
# --------------------------------------------------------------------------


def build(output_dir: Path) -> dict[str, Any]:
    targets = {t["case_id"]: t for t in build_targets()}
    qna = {int(r["id"]): r for r in json.loads((BASELINE / "qna-canonical.json").read_text(encoding="utf-8")) if r["status"] == 1}
    guards = {int(r["qna_id"]): r for r in json.loads((BASELINE / "qna-routing-guards.json").read_text(encoding="utf-8"))}
    alias_owners: dict[str, set[int]] = defaultdict(set)
    for row in json.loads((BASELINE / "qna-aliases.json").read_text(encoding="utf-8")):
        alias_owners[row["query_text"].strip()].add(int(row["qna_id"]))

    views: dict[str, dict[int, dict[str, Any]]] = {}
    for model in MODELS:
        results = {r["case_id"]: r for r in read_jsonl(BENCH / model / "results.jsonl")}
        views[model] = {}
        for case_id, result in results.items():
            target = targets[case_id]
            evaluation = evaluate(target, result, set(guards), alias_owners)
            if target["scorable"]:
                views[model][case_id] = trace_view(result, evaluation, set(target["expected_qna_ids"]))

    scorable = sorted(views[MODELS[0]])
    cases, false_splits, calendar_rows, kb_rows = [], [], [], []
    for case_id in scorable:
        target = targets[case_id]
        a, b = views["4o-mini"][case_id], views["luna-high"][case_id]
        exact_bucket = ("both_correct" if a["exact_correct"] and b["exact_correct"] else "only_4o_mini" if a["exact_correct"]
                        else "only_luna_high" if b["exact_correct"] else "both_wrong")
        relaxed_bucket = ("both_correct" if a["relaxed_correct"] and b["relaxed_correct"] else "only_4o_mini" if a["relaxed_correct"]
                          else "only_luna_high" if b["relaxed_correct"] else "both_wrong")
        a_cause, a_secondary = fine_cause(target, a, qna, alias_owners, b)
        b_cause, b_secondary = fine_cause(target, b, qna, alias_owners, a)
        record = {
            "case_id": case_id, "user_message": target["user_message"], "turn_type": target["turn_type"],
            "expected_intent_groups": target["expected_intent_groups"], "expected_qna_ids": target["expected_qna_ids"],
            "gold": [{"qna_id": g, "question": qna[g]["question_text"], "answer": qna[g]["answer_text"][:600]}
                     for g in target["expected_qna_ids"] if g in qna],
            "retrieval": {"first_pool_gold_qna_rank": a["speculative_gold_rank"]["qna_rank"],
                          "first_pool_gold_prompt_index": a["speculative_gold_rank"]["prompt_index"],
                          "calendar_candidates": a["speculative_gold_rank"].get("calendar_candidates"),
                          "same_first_pool_both_models": a["speculative_candidates"] == b["speculative_candidates"]},
            "exact_bucket": exact_bucket, "relaxed_bucket": relaxed_bucket,
            "4o_mini": {**a, "fine_cause": a_cause, "secondary_causes": a_secondary},
            "luna_high": {**b, "fine_cause": b_cause, "secondary_causes": b_secondary},
            "exact_alias_owners": sorted(alias_owners.get(target["user_message"].strip(), set())),
            "gold_guarded": bool(set(target["expected_qna_ids"]) & set(guards)),
        }
        if exact_bucket in ("only_4o_mini", "only_luna_high"):
            record["model_diff_labels"] = model_diff_labels(a, b)
        record["needs_human_review"] = any(c == "NEEDS_HUMAN_REVIEW" for c in (a_cause, b_cause)) or any(
            "GOLD_REVIEW_CANDIDATE" in s for s in (a_secondary, b_secondary))
        if exact_bucket != "both_correct" or a["split_count"] > 1 or b["split_count"] > 1 or a["calendar_selected"] or b["calendar_selected"]:
            cases.append(record)

        for label, view in (("4o-mini", a), ("luna-high", b)):
            if len(target["expected_intent_groups"]) == 1 and view["split_count"] > 1:
                extra = [s for s in view["selected_qna_ids"] if s not in target["expected_qna_ids"]]
                if view["exact_correct"]:
                    kind = "HARMLESS_FALSE_SPLIT"
                elif view["relaxed_correct"]:
                    kind = "EXTRA_ANSWER_FALSE_SPLIT"
                elif view["speculative_correct"]:
                    kind = "CONTEXT_LOSS_FALSE_SPLIT"
                else:
                    kind = "WRONG_ANSWER_FALSE_SPLIT"
                gold_len = sum(len(qna[g]["answer_text"]) for g in target["expected_qna_ids"] if g in qna)
                false_splits.append({
                    "model": label, "case_id": case_id, "class": kind, "pieces": view["split_count"],
                    "sub_questions": view["sub_questions"],
                    "fragments_selected_same_qna": len({a2["selected_qna_id"] for a2 in view["used_asks"]}) == 1,
                    "extra_answer_ids": extra, "gold_still_returned": view["relaxed_correct"],
                    "exact_broken": not view["exact_correct"], "speculative_full_message_correct": view["speculative_correct"],
                    "final_answer_chars": view["final_answer_chars"], "gold_answer_chars": gold_len,
                    "answer_length_ratio": round(view["final_answer_chars"] / gold_len, 2) if gold_len else None,
                    "independence": "NEEDS_HUMAN_REVIEW",
                })
            if view["calendar_selected"]:
                calendar_ask = next((u for u in view["used_asks"] if u["selected_calendar"]), None)
                date_signal = bool(DATE_SIGNAL.search(target["user_message"]))
                gold_rank = view["speculative_gold_rank"]["qna_rank"]
                if view["relaxed_correct"]:
                    kind = "CALENDAR_NOT_PRIMARY_CAUSE"
                elif date_signal:
                    kind = "CALENDAR_AMBIGUOUS_DATE_SIGNAL"
                elif gold_rank and gold_rank <= 3:
                    kind = "CALENDAR_SELECTED_DESPITE_CLEAR_QNA"
                else:
                    kind = "CALENDAR_HARMFUL_DISTRACTOR"
                calendar_rows.append({
                    "model": label, "case_id": case_id, "class": kind, "user_message": target["user_message"],
                    "gold": [qna[g]["question_text"] for g in target["expected_qna_ids"] if g in qna],
                    "calendar_candidate": calendar_ask["selected_question"] if calendar_ask else None,
                    "date_signal_in_message": date_signal, "gold_qna_rank": gold_rank,
                    "relaxed_correct": view["relaxed_correct"],
                    "other_model_correct": (b if label == "4o-mini" else a)["exact_correct"],
                    "speculative_full_message_correct": view["speculative_correct"],
                })

    # KB / yakın-kopya analizi (seçicinin gold havuzdayken yanlış seçtiği)
    pair_counter: Counter = Counter()
    for record in cases:
        for key in ("4o_mini", "luna_high"):
            view = record[key]
            if view["fine_cause"] in ("SELECTOR_NEAR_DUPLICATE_CONFUSION", "SELECTOR_WRONG_SEMANTIC_MATCH", "KB_OVERLAP_AMBIGUITY"):
                wrong = [s for s in view["selected_qna_ids"] if s not in record["expected_qna_ids"]]
                for chosen in wrong[:1]:
                    gold_id = record["expected_qna_ids"][0]
                    pair_counter[(gold_id, chosen)] += 1
                    kb_rows.append({
                        "model": key, "case_id": record["case_id"], "fine_cause": view["fine_cause"],
                        "user_message": record["user_message"],
                        "gold_qna_id": gold_id, "gold_question": qna[gold_id]["question_text"],
                        "selected_qna_id": chosen, "selected_question": qna.get(chosen, {}).get("question_text"),
                        "question_similarity": similarity(qna[gold_id]["question_text"], qna.get(chosen, {}).get("question_text", "")),
                        "answer_similarity": similarity(qna[gold_id]["answer_text"], qna.get(chosen, {}).get("answer_text", "")),
                        "message_is_alias_of_selected": chosen in alias_owners.get(record["user_message"].strip(), set()),
                        "message_is_alias_of_gold": gold_id in alias_owners.get(record["user_message"].strip(), set()),
                        "both_models_same_wrong": "BOTH_MODELS_SAME_WRONG_CHOICE" in view["secondary_causes"],
                    })
    recurring_pairs = [{"gold_qna_id": g, "selected_qna_id": s, "count": n,
                        "gold_question": qna[g]["question_text"], "selected_question": qna.get(s, {}).get("question_text")}
                       for (g, s), n in pair_counter.most_common() if n >= 2]

    report = summarize(cases, views, targets, guards, qna, false_splits, calendar_rows, kb_rows, recurring_pairs)
    output_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "paired-error-cases.jsonl": jsonl(cases),
        "false-split-analysis.jsonl": jsonl(false_splits),
        "calendar-analysis.jsonl": jsonl(calendar_rows),
        "kb-ambiguity-analysis.jsonl": jsonl(kb_rows),
        "root-cause-summary.json": dump(report["root_causes"]),
        "model-behavior-comparison.json": dump(report["behavior"]),
        "multi-intent-analysis.json": dump(report["multi_intent"]),
        "guarded-analysis.json": dump(report["guarded"]),
        "position-bias-analysis.json": dump(report["position_bias"]),
        "paired-error-audit.json": dump(report),
        "PAIRED-ERROR-AUDIT.md": render(report),
    }
    for name, content in files.items():
        (output_dir / name).write_text(content, encoding="utf-8")
    build_workbook(cases, false_splits, calendar_rows, report, output_dir / "paired-error-human-review.xlsx")
    report["outputs_sha256"] = {name: sha256(output_dir / name) for name in sorted(files) if name != "paired-error-audit.json"}
    return {"report": report, "cases": cases, "false_splits": false_splits, "calendar": calendar_rows, "kb": kb_rows}


def summarize(cases, views, targets, guards, qna, false_splits, calendar_rows, kb_rows, recurring_pairs) -> dict[str, Any]:
    buckets = Counter(c["exact_bucket"] for c in cases if c["exact_bucket"] != "both_correct")
    relaxed = Counter(c["relaxed_bucket"] for c in cases)
    universe = [c for c in cases if c["exact_bucket"] != "both_correct"]
    extra_answer = [c for c in cases if any(c[k]["relaxed_correct"] and not c[k]["exact_correct"] for k in ("4o_mini", "luna_high"))]

    def causes_for(bucket: str, key: str) -> dict[str, int]:
        return dict(Counter(c[key]["fine_cause"] for c in universe if c["exact_bucket"] == bucket and c[key]["fine_cause"]).most_common())

    diff_labels = Counter(label for c in universe for label in c.get("model_diff_labels", []))
    pure = [c["case_id"] for c in universe if "PURE_SELECTOR_MODEL_DIFFERENCE" in c.get("model_diff_labels", [])]
    pure_split = Counter("4o_mini_correct" if c["4o_mini"]["exact_correct"] else "luna_correct" for c in universe
                         if c["case_id"] in pure)

    all_scorable = sorted(views["4o-mini"])
    behavior = {}
    for model, key in (("4o-mini", "4o_mini"), ("luna-high", "luna_high")):
        view = views[model]
        single = [c for c in all_scorable if len(targets[c]["expected_intent_groups"]) == 1]
        gold_in_spec = [c for c in single if view[c]["speculative_gold_rank"]["qna_rank"]]
        fine = Counter(c[key]["fine_cause"] for c in cases if c[key]["fine_cause"])
        behavior[model] = {
            "false_split_rate": {"n": sum(1 for c in single if view[c]["split_count"] > 1), "of": len(single)},
            "none_rate_full_message": {"n": sum(1 for c in all_scorable if view[c]["speculative_declined"]), "of": len(all_scorable)},
            "none_when_gold_in_pool_full_message": sum(1 for c in gold_in_spec if view[c]["speculative_declined"]),
            "full_message_selector_accuracy": {"n": sum(1 for c in gold_in_spec if view[c]["speculative_correct"]),
                                               "of": len(gold_in_spec),
                                               "note": "Aynı tam mesaj + aynı aday havuzu; bölücüden bağımsız"},
            "wrong_positive_full_message": sum(1 for c in gold_in_spec if not view[c]["speculative_correct"]
                                               and not view[c]["speculative_declined"]),
            "calendar_selected_full_message": sum(1 for c in all_scorable if view[c]["speculative_calendar"]),
            "calendar_selected_final": sum(1 for c in all_scorable if view[c]["calendar_selected"]),
            "near_duplicate_confusion": fine.get("SELECTOR_NEAR_DUPLICATE_CONFUSION", 0),
            "kb_overlap_ambiguity": fine.get("KB_OVERLAP_AMBIGUITY", 0),
            "exact_accuracy": {"n": sum(1 for c in all_scorable if view[c]["exact_correct"]), "of": len(all_scorable)},
            "relaxed_accuracy": {"n": sum(1 for c in all_scorable if view[c]["relaxed_correct"]), "of": len(all_scorable)},
            "fine_cause_counts": dict(fine.most_common()),
        }
    # Aynı tam-mesaj girdisinde eşli seçici karşılaştırması
    same_input = [c for c in all_scorable if views["4o-mini"][c]["speculative_candidates"] == views["luna-high"][c]["speculative_candidates"]
                  and views["4o-mini"][c]["speculative_gold_rank"]["qna_rank"]]
    spec_pairs = Counter(("A" if views["4o-mini"][c]["speculative_correct"] else "a") + ("B" if views["luna-high"][c]["speculative_correct"] else "b")
                         for c in same_input)
    from scripts.run_e2e_baseline import mcnemar

    behavior["full_message_paired"] = {"cases": len(same_input), "both_correct": spec_pairs["AB"], "only_4o_mini": spec_pairs["Ab"],
                                       "only_luna_high": spec_pairs["aB"], "both_wrong": spec_pairs["ab"],
                                       "mcnemar": mcnemar(spec_pairs["Ab"], spec_pairs["aB"])}

    # Pozisyon analizi (yalnız bölünmemiş, gold havuzda)
    position: dict[str, Any] = {}
    for model in MODELS:
        view = views[model]
        buckets_rank: dict[str, list[bool]] = defaultdict(list)
        wrong_ranks = []
        for c in all_scorable:
            v = view[c]
            rank = v["speculative_gold_rank"]["qna_rank"]
            # Tam mesaj seçimi iki modelde aynı girdi ve aday sırasıyla yapılır; bölmeden bağımsızdır.
            if not rank or len(targets[c]["expected_intent_groups"]) != 1:
                continue
            bucket = "1" if rank == 1 else "2-3" if rank <= 3 else "4-5" if rank <= 5 else "6+"
            buckets_rank[bucket].append(v["speculative_correct"])
            selected_rank = v["speculative_gold_rank"].get("selected_qna_rank")
            if not v["speculative_correct"] and selected_rank:
                wrong_ranks.append({"case_id": c, "gold_rank": rank, "selected_rank": selected_rank, "delta": selected_rank - rank})
        position[model] = {
            "accuracy_by_gold_rank": {k: {"n": sum(v), "of": len(v), "rate": round(sum(v) / len(v), 4)}
                                      for k, v in sorted(buckets_rank.items(), key=lambda kv: ["1", "2-3", "4-5", "6+"].index(kv[0]))},
            "wrong_selection_rank_distribution": dict(sorted(Counter(w["selected_rank"] for w in wrong_ranks).items())),
            "wrong_selected_above_gold": sum(1 for w in wrong_ranks if w["delta"] < 0),
            "wrong_selected_below_gold": sum(1 for w in wrong_ranks if w["delta"] > 0),
            "wrong_cases": len(wrong_ranks),
        }

    # Guard'lı hedefler
    guarded_cases = [c for c in all_scorable if set(targets[c]["expected_qna_ids"]) & set(guards)]
    guarded = {"targets": len(guarded_cases), "cases": []}
    for c in guarded_cases:
        a, b = views["4o-mini"][c], views["luna-high"][c]
        gold = targets[c]["expected_qna_ids"]
        guarded["cases"].append({
            "case_id": c, "gold": gold, "guard_modes": sorted({guards[g]["content_mode"] for g in gold if g in guards}),
            "gold_question": [qna[g]["question_text"] for g in gold if g in qna],
            "same_candidates": a["speculative_candidates"] == b["speculative_candidates"],
            "gold_rank": a["speculative_gold_rank"]["qna_rank"],
            "4o_mini": {"exact": a["exact_correct"], "relaxed": a["relaxed_correct"], "selected": a["selected_qna_ids"],
                        "split": a["split_count"], "declined": a["declined_any"]},
            "luna_high": {"exact": b["exact_correct"], "relaxed": b["relaxed_correct"], "selected": b["selected_qna_ids"],
                          "split": b["split_count"], "declined": b["declined_any"],
                          "selected_questions": [qna[s]["question_text"] for s in b["selected_qna_ids"] if s in qna]},
        })
    for model, key in (("4o-mini", "4o_mini"), ("luna-high", "luna_high")):
        guarded[model] = {"exact": sum(1 for g in guarded["cases"] if g[key]["exact"]),
                          "relaxed": sum(1 for g in guarded["cases"] if g[key]["relaxed"]),
                          "declined": sum(1 for g in guarded["cases"] if g[key]["declined"])}
    guarded["same_candidates_all"] = all(g["same_candidates"] for g in guarded["cases"])
    guarded["by_mode"] = {}
    for g in guarded["cases"]:
        for mode in g["guard_modes"]:
            entry = guarded["by_mode"].setdefault(mode, {"targets": 0, "4o_mini_exact": 0, "luna_high_exact": 0})
            entry["targets"] += 1
            entry["4o_mini_exact"] += g["4o_mini"]["exact"]
            entry["luna_high_exact"] += g["luna_high"]["exact"]

    multi = {}
    for c in (71, 480):
        multi[str(c)] = {"expected_intent_groups": targets[c]["expected_intent_groups"], "user_message": targets[c]["user_message"]}
        for model in MODELS:
            v = views[model][c]
            groups = targets[c]["expected_intent_groups"]
            per_fragment = []
            for ask in v["used_asks"] or [None]:
                if ask is None:
                    continue
                per_fragment.append({"fragment": ask["question"], "selected": ask["selected_qna_id"],
                                     "declined": ask["declined"],
                                     "groups_in_pool": [bool(set(g) & set(q for q in ask["candidate_ids"] if q)) for g in groups]})
            lost = []
            for index, group in enumerate(groups):
                if set(group) & set(v["selected_qna_ids"]):
                    continue
                in_any = any(f["groups_in_pool"][index] for f in per_fragment)
                lost.append({"group": group, "gold_in_some_fragment_pool": in_any,
                             "loss_stage": "SELECTOR" if in_any else ("SPLITTER" if v["split_count"] < len(groups) else "RETRIEVAL")})
            multi[str(c)][model] = {"split_count": v["split_count"], "sub_questions": v["sub_questions"],
                                    "selected": v["selected_qna_ids"], "fragments": per_fragment, "lost_groups": lost,
                                    "exact": v["exact_correct"], "relaxed": v["relaxed_correct"]}

    none_calibration = {}
    for model in MODELS:
        view = views[model]
        declined = [c for c in all_scorable if view[c]["speculative_declined"]]
        none_calibration[model] = {
            "declined_full_message": len(declined),
            "declined_with_gold_in_pool": sum(1 for c in declined if view[c]["speculative_gold_rank"]["qna_rank"]),
            "declined_without_gold_in_pool": sum(1 for c in declined if not view[c]["speculative_gold_rank"]["qna_rank"]),
            "declined_then_final_exact_correct": sum(1 for c in declined if view[c]["exact_correct"]),
            "declined_then_fallback_used": sum(1 for c in declined if view[c]["fallback"]),
        }

    fs = Counter((r["model"], r["class"]) for r in false_splits)
    false_split_summary = {model: {cls: fs.get((model, cls), 0) for cls in
                                   ("HARMLESS_FALSE_SPLIT", "EXTRA_ANSWER_FALSE_SPLIT", "CONTEXT_LOSS_FALSE_SPLIT", "WRONG_ANSWER_FALSE_SPLIT")}
                           for model in MODELS}
    for model in MODELS:
        rows = [r for r in false_splits if r["model"] == model]
        false_split_summary[model]["total"] = len(rows)
        false_split_summary[model]["exact_broken"] = sum(r["exact_broken"] for r in rows)
        false_split_summary[model]["mean_answer_length_ratio"] = round(
            sum(r["answer_length_ratio"] or 0 for r in rows) / len(rows), 2) if rows else None
        false_split_summary[model]["pieces_distribution"] = dict(sorted(Counter(r["pieces"] for r in rows).items()))

    calendar_summary = {model: dict(Counter(r["class"] for r in calendar_rows if r["model"] == model)) for model in MODELS}

    decision_rows = decision_table(behavior, false_split_summary, calendar_summary, position, none_calibration, cases, targets)
    human_review = sorted(c["case_id"] for c in cases if c["needs_human_review"])
    return {
        "universe": {
            "metric": "exact (birebir) — fazladan cevap hata sayılır",
            "only_4o_mini_correct": buckets.get("only_4o_mini", 0), "only_luna_high_correct": buckets.get("only_luna_high", 0),
            "both_wrong": buckets.get("both_wrong", 0), "total": len(universe),
            "relaxed_buckets_over_listed_cases": dict(relaxed),
            "extra_answer_cases": len(extra_answer),
            "note": ("Kullanıcı özetindeki 64 'ikisi de yanlış' gevşek metriğe aittir; birebir metrikte 75. "
                     "Evren tutarlılık için birebir metrikle kuruldu; gevşek sonuç her vakada ayrı alan."),
        },
        "root_causes": {"only_4o_mini_correct": {"luna_high_cause": causes_for("only_4o_mini", "luna_high")},
                        "only_luna_high_correct": {"4o_mini_cause": causes_for("only_luna_high", "4o_mini")},
                        "both_wrong": {"4o_mini_cause": causes_for("both_wrong", "4o_mini"),
                                       "luna_high_cause": causes_for("both_wrong", "luna_high")},
                        "model_diff_labels": dict(diff_labels.most_common()),
                        "pure_selector_model_difference": {"cases": len(pure), "split": dict(pure_split), "case_ids": pure}},
        "behavior": behavior,
        "false_split": false_split_summary,
        "calendar": calendar_summary,
        "multi_intent": multi,
        "guarded": guarded,
        "position_bias": position,
        "none_calibration": none_calibration,
        "kb_ambiguity": {"rows": len(kb_rows), "recurring_gold_selected_pairs": recurring_pairs,
                         "message_is_alias_of_selected": sum(1 for r in kb_rows if r["message_is_alias_of_selected"]),
                         "near_duplicate_rows": sum(1 for r in kb_rows if r["fine_cause"] == "SELECTOR_NEAR_DUPLICATE_CONFUSION")},
        "context": {"context_required_unscored": sum(1 for t in targets.values() if not t["scorable"]),
                    "context_available_failures": {m: sum(1 for c in cases if c["turn_type"] == "FOLLOW_UP_CONTEXT_AVAILABLE_BUT_NOT_REQUIRED"
                                                          and not c[k]["exact_correct"])
                                                   for m, k in (("4o-mini", "4o_mini"), ("luna-high", "luna_high"))}},
        "decision_table": decision_rows,
        "human_review": {"cases": len(human_review), "case_ids": human_review},
        "sources": {name: sha256(BENCH / name) for name in ("benchmark-manifest.json", "comparison.json", "paired-cases.jsonl")},
    }


def decision_table(behavior, false_split, calendar, position, none_calibration, cases, targets) -> list[dict[str, Any]]:
    def count(model_key: str, causes: tuple[str, ...]) -> int:
        return sum(1 for c in cases if c[model_key]["fine_cause"] in causes)

    rows = []
    for problem, causes, shared_note, intervention in [
        ("Gereksiz bölme (fazladan cevap / bağlam kaybı)", ("FALSE_SPLIT_SINGLE_INTENT", "FRAGMENT_LOST_CONTEXT", "SELECTOR_FRAGMENT_CONTEXT_LOSS"),
         "Bölücü davranışı modele bağlı; birleştirme mimarisi ortak", "Bölme kararı/eşiği; spekülatif sonucu tercih"),
        ("Seçici yakın-kopya karışıklığı", ("SELECTOR_NEAR_DUPLICATE_CONFUSION",), "Ortak (KB'de yakın QnA'lar)", "KB ayrıştırma / aday açıklaması"),
        ("Seçici anlamsal yanlış eşleşme", ("SELECTOR_WRONG_SEMANTIC_MATCH",), "Modele bağlı", "Seçici prompt / model"),
        ("KB örtüşmesi (mesaj seçilen QnA'nın alias'ı)", ("KB_OVERLAP_AMBIGUITY",), "Ortak (KB/gold)", "Gold/KB inceleme"),
        ("None kalibrasyonu (gold varken 'hiçbiri')", ("SELECTOR_NONE_WHEN_GOLD_PRESENT",), "Modele bağlı", "Seçici eşik/prompt"),
        ("Takvim dikkat dağıtıcı", ("SELECTOR_CALENDAR_DISTRACTION", "CALENDAR_ROUTING_INTERFERENCE"), "Ortak mimari (takvim havuzun başında)", "Takvim yönlendirmesi"),
        ("Retrieval kaçırma", ("RETRIEVAL_TRUE_MISS",), "Ortak mimari", "Retrieval"),
        ("Çoklu niyet kaçırma", ("MISSED_MULTI_INTENT",), "Modele bağlı", "Bölücü"),
    ]:
        rows.append({"problem": problem, "4o_mini": count("4o_mini", causes), "luna_high": count("luna_high", causes),
                     "shared_architecture": shared_note, "potential_intervention": intervention,
                     "exact_accuracy_upper_bound_gain": {"4o_mini": count("4o_mini", causes), "luna_high": count("luna_high", causes)}})
    rows.append({"problem": "Pozisyon yanlılığı sinyali (yanlış seçim gold'un üstünde, gold sırası ≥4)",
                 "4o_mini": sum(1 for c in cases if "SELECTOR_POSITION_BIAS_SIGNAL" in c["4o_mini"]["secondary_causes"]),
                 "luna_high": sum(1 for c in cases if "SELECTOR_POSITION_BIAS_SIGNAL" in c["luna_high"]["secondary_causes"]),
                 "shared_architecture": "Aday sırası ortak", "potential_intervention": "Aday sıralama/karıştırma",
                 "exact_accuracy_upper_bound_gain": None})
    rows.append({"problem": "Bağlam (puanlanamadı)", "4o_mini": None, "luna_high": None,
                 "shared_architecture": "Gold eksik (17 hedef)", "potential_intervention": "Bağlam hedeflerine gold",
                 "exact_accuracy_upper_bound_gain": None})
    return rows


def render(report: dict[str, Any]) -> str:
    u = report["universe"]
    lines = ["# Eşli model hata denetimi v1", "",
             f"Evren (birebir metrik): yalnız 4o-mini doğru {u['only_4o_mini_correct']} · yalnız Luna doğru {u['only_luna_high_correct']} · "
             f"ikisi de yanlış {u['both_wrong']} = **{u['total']}**", "", f"> {u['note']}", "",
             "## Karar tablosu", "", "| Problem | 4o-mini | Luna | Ortak mimari mi? | Potansiyel müdahale |", "|---|---:|---:|---|---|"]
    lines += [f"| {r['problem']} | {r['4o_mini'] if r['4o_mini'] is not None else '—'} | {r['luna_high'] if r['luna_high'] is not None else '—'} | "
              f"{r['shared_architecture']} | {r['potential_intervention']} |" for r in report["decision_table"]]
    b = report["behavior"]
    lines += ["", "## Tam mesaj seçici (aynı girdi, bölücüden bağımsız)", "",
              f"- 4o-mini: {b['4o-mini']['full_message_selector_accuracy']['n']}/{b['4o-mini']['full_message_selector_accuracy']['of']}",
              f"- Luna-high: {b['luna-high']['full_message_selector_accuracy']['n']}/{b['luna-high']['full_message_selector_accuracy']['of']}",
              f"- Eşli: `{json.dumps(b['full_message_paired'], ensure_ascii=False)}`", "",
              "## Gereksiz bölme", "", f"`{json.dumps(report['false_split'], ensure_ascii=False)}`", "",
              "## Saf seçici model farkı", "", f"`{json.dumps(report['root_causes']['pure_selector_model_difference'], ensure_ascii=False)[:600]}`", "",
              "## İnsan incelemesi", "", f"- {report['human_review']['cases']} vaka", ""]
    return "\n".join(lines)


def build_workbook(cases, false_splits, calendar_rows, report, path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Özet"
    u = report["universe"]
    for row in (["Eşli model hata denetimi v1"], [], ["Toplam denetim vakası (birebir)", u["total"]],
                ["Yalnız 4o-mini doğru", u["only_4o_mini_correct"]], ["Yalnız Luna doğru", u["only_luna_high_correct"]],
                ["İkisi de yanlış", u["both_wrong"]], ["Otomatik sınıflanan", u["total"] - report["human_review"]["cases"]],
                ["İnsan incelemesi gereken", report["human_review"]["cases"]], [], ["Not", u["note"]]):
        summary.append(row)
    summary["A1"].font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 100

    headers = ["Vaka", "Reviewer decision", "Reviewer note", "Mesaj", "Turn", "Gold QnA", "Gold soru", "Gold sırası",
               "4o-mini seçim", "4o-mini bölme", "4o-mini birebir", "4o-mini neden", "Luna seçim", "Luna bölme",
               "Luna birebir", "Luna neden", "Model farkı", "İkincil (4o / Luna)"]

    def case_sheet(title: str, rows: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = HEADER
            cell.font = Font(bold=True)
        for index, c in enumerate(rows, start=2):
            a, b = c["4o_mini"], c["luna_high"]
            sheet.append([c["case_id"], None, None, c["user_message"], c["turn_type"], ", ".join(map(str, c["expected_qna_ids"])),
                          " / ".join(g["question"] for g in c["gold"]), c["retrieval"]["first_pool_gold_qna_rank"],
                          ", ".join(map(str, a["selected_qna_ids"])) or ("KATILMADI" if a["declined_any"] else "—"),
                          a["split_count"], a["exact_correct"], a["fine_cause"] or "—",
                          ", ".join(map(str, b["selected_qna_ids"])) or ("KATILMADI" if b["declined_any"] else "—"),
                          b["split_count"], b["exact_correct"], b["fine_cause"] or "—",
                          ", ".join(c.get("model_diff_labels", [])),
                          f"{', '.join(a['secondary_causes'])} / {', '.join(b['secondary_causes'])}"])
            sheet[f"B{index}"].fill = YELLOW
            sheet[f"C{index}"].fill = YELLOW
        validation = DataValidation(type="list", formula1='"' + ",".join(REVIEW_DECISIONS) + '"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"B2:B{max(len(rows) + 1, 2)}")
        for position, width in enumerate([7, 22, 30, 60, 22, 12, 50, 8, 16, 8, 9, 30, 16, 8, 9, 30, 40, 40], start=1):
            sheet.column_dimensions[get_column_letter(position)].width = width
        sheet.freeze_panes = "D2"

    universe = [c for c in cases if c["exact_bucket"] != "both_correct"]
    case_sheet("Only 4o Correct", [c for c in universe if c["exact_bucket"] == "only_4o_mini"])
    case_sheet("Only Luna Correct", [c for c in universe if c["exact_bucket"] == "only_luna_high"])
    case_sheet("Both Wrong", [c for c in universe if c["exact_bucket"] == "both_wrong"])

    sheet = workbook.create_sheet("False Splits")
    sheet.append(["Model", "Vaka", "Sınıf", "Parça", "Alt sorular", "Fazladan cevap", "Gold döndü", "Birebir bozuldu",
                  "Tam mesaj doğruydu", "Cevap uzunluk oranı", "Parçalar bağımsız mı (insan)"])
    for r in false_splits:
        sheet.append([r["model"], r["case_id"], r["class"], r["pieces"], " | ".join(r["sub_questions"]),
                      ", ".join(map(str, r["extra_answer_ids"])), r["gold_still_returned"], r["exact_broken"],
                      r["speculative_full_message_correct"], r["answer_length_ratio"], None])
    sheet = workbook.create_sheet("Calendar")
    sheet.append(["Model", "Vaka", "Sınıf", "Mesaj", "Gold", "Takvim adayı", "Tarih sinyali", "Gold sırası", "Diğer model doğru"])
    for r in calendar_rows:
        sheet.append([r["model"], r["case_id"], r["class"], r["user_message"], " / ".join(r["gold"]), r["calendar_candidate"],
                      r["date_signal_in_message"], r["gold_qna_rank"], r["other_model_correct"]])
    sheet = workbook.create_sheet("Multi Intent")
    sheet.append(["Vaka", "Model", "Bölme", "Alt sorular", "Seçilen", "Kaybolan grup", "Kayıp aşaması"])
    for case_id, entry in report["multi_intent"].items():
        for model in MODELS:
            m = entry[model]
            sheet.append([int(case_id), model, m["split_count"], " | ".join(m["sub_questions"]), ", ".join(map(str, m["selected"])),
                          "; ".join(str(g["group"]) for g in m["lost_groups"]), "; ".join(g["loss_stage"] for g in m["lost_groups"])])
    case_sheet("Needs Human Review", [c for c in cases if c["needs_human_review"]])
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs" / "paired-error-audit-v1-20260918")
    options = parser.parse_args()
    result = build(options.output_dir)
    report = result["report"]
    (options.output_dir / "paired-error-audit.json").write_text(dump(report), encoding="utf-8")
    print(json.dumps({"universe": report["universe"], "human_review": report["human_review"]["cases"],
                      "pure_selector": report["root_causes"]["pure_selector_model_difference"]["cases"]},
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
