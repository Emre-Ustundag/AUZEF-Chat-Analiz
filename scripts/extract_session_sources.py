"""Gold vakalarının değdiği gerçek oturumları ham ``chatbot.xlsx``'ten çıkarır.

Tek akışlı, deterministik, LLM yok. ``prepare_sessions.py`` ile AYNI kullanıcı
turn filtresi uygulanır; böylece çıkarılan kullanıcı turn sırası
``sessions-yil.jsonl`` ve ``alias-session-matches.jsonl`` turn_index'leriyle
birebir örtüşür. Bot mesajları da (bağlam olarak) korunur; metinler mevcut
hattaki gibi PII maskelemesinden geçer.

Çıktı: oturum başına bir JSON satırı (mesaj kimliği, sıra, yön, zaman, metin).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "apps" / "backend"))

from app.pipeline.preprocess import COURTESY_ONLY, is_system_message, normalize  # noqa: E402
from app.services.redaction import redact_pii  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

WS = re.compile(r"\s+")
FALLBACK = "sizi ne yazık ki anlayamadım"
MIN_LEN = 3
USER_DIRECTION = "Kullanıcı"


def file_digest(path: Path) -> str:
    digest = hashlib.blake2b(digest_size=32)
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_gold_user_turn(text: str, message_type: str, quick_reply: str) -> bool:
    """prepare_sessions.py'deki kullanıcı turn filtresinin birebir aynısı."""
    if quick_reply in {"Onayladı", "Reddetti"}:
        return False
    if message_type != "text":
        return False
    return not (
        not text
        or text == "None"
        or len(text) < MIN_LEN
        or is_system_message(text)
        or normalize(text) in COURTESY_ONLY
    )


def wanted_sessions(alias_matches: Path) -> set[str]:
    sessions: set[str] = set()
    for line in alias_matches.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        sessions.update(o["session_id"] for o in row.get("occurrences") or [])
        sessions.update(row.get("matched_session_ids") or [])
    return sessions


def extract(source: Path, alias_matches: Path, output_dir: Path) -> dict[str, Any]:
    wanted = wanted_sessions(alias_matches)
    workbook = load_workbook(source, read_only=True)
    sessions: dict[str, dict[str, Any]] = {}
    scanned = 0
    try:
        for sheet_name in workbook.sheetnames:
            rows = workbook[sheet_name].iter_rows(values_only=True)
            index = {header: position for position, header in enumerate(next(rows))}
            for row in rows:
                if row is None:
                    continue
                scanned += 1
                session_id = row[index["SessionId"]]
                if not session_id or str(session_id) not in wanted:
                    continue
                session_id = str(session_id)
                session = sessions.setdefault(session_id, {
                    "session_id": session_id,
                    "channel": str(row[index["Channel"]] or ""),
                    "session_status": str(row[index["SessionStatus"]] or ""),
                    "session_start": str(row[index["SessionStartTr"]] or ""),
                    "session_end": str(row[index["SessionEndTr"]] or ""),
                    "duration_minutes": row[index["DurationMinutes"]],
                    "source_sheets": [],
                    "messages": [],
                })
                if sheet_name not in session["source_sheets"]:
                    session["source_sheets"].append(sheet_name)
                text = WS.sub(" ", str(row[index["MessageTextClean"]] or "")).strip()
                message_type = str(row[index["MessageType"]] or "")
                quick_reply = str(row[index["QuickReplyLabel"]] or "").strip()
                direction = str(row[index["Direction"]] or "")
                session["messages"].append({
                    "message_id": int(row[index["MessageId"]]) if row[index["MessageId"]] is not None else None,
                    "message_order": int(row[index["MessageOrder"]]) if row[index["MessageOrder"]] is not None else None,
                    "direction": direction,
                    "message_type": message_type,
                    "quick_reply": quick_reply,
                    "time": str(row[index["MessageTimeTr"]] or ""),
                    "text": redact_pii(text),
                    "is_gold_user_turn": direction == USER_DIRECTION and is_gold_user_turn(text, message_type, quick_reply),
                    "is_bot_fallback": direction != USER_DIRECTION and text.casefold().startswith(FALLBACK),
                })
    finally:
        workbook.close()

    output_dir.mkdir(parents=True, exist_ok=True)
    extract_path = output_dir / "sessions-extract.jsonl"
    with extract_path.open("w", encoding="utf-8") as handle:
        for session_id in sorted(sessions):
            session = sessions[session_id]
            session["messages"].sort(key=lambda m: (m["message_order"] if m["message_order"] is not None else 0,
                                                    m["message_id"] if m["message_id"] is not None else 0))
            turn = 0
            for message in session["messages"]:
                if message["is_gold_user_turn"]:
                    turn += 1
                    message["user_turn_index"] = turn
                else:
                    message["user_turn_index"] = None
            session["user_turn_count"] = turn
            handle.write(json.dumps(session, ensure_ascii=False, sort_keys=True) + "\n")

    report = {
        "source": {"name": source.name, "blake2b": file_digest(source), "rows_scanned": scanned},
        "alias_matches": {"name": alias_matches.name, "blake2b": file_digest(alias_matches)},
        "requested_sessions": len(wanted),
        "extracted_sessions": len(sessions),
        "missing_sessions": sorted(wanted - set(sessions)),
        "messages": sum(len(s["messages"]) for s in sessions.values()),
        "user_turns": sum(s["user_turn_count"] for s in sessions.values()),
        "extract_file": extract_path.name,
        "extract_blake2b": file_digest(extract_path),
    }
    (output_dir / "extract-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="Ham chatbot.xlsx")
    parser.add_argument("--alias-matches", type=Path, default=ROOT / "outputs" / "gold-v2-sources-20260917" / "alias-session-matches.jsonl")
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    options = parse_args()
    report = extract(options.source, options.alias_matches, options.output_dir)
    print(json.dumps({k: v for k, v in report.items() if k != "missing_sessions"}, ensure_ascii=False, indent=2))
    print(f"missing_sessions: {len(report['missing_sessions'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
