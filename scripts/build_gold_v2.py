"""İnsan karar workbook'unu doğrular ve hazırsa makine-okunur Gold v2 üretir.

Workbook hiçbir zaman değiştirilmez. Doğrulama raporu her koşuda yazılır;
Gold v2 yalnız bütün zorunlu insan alanları geçerliyse oluşturulur.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DECISIONS = {
    "4o daha iyi",
    "Luna daha iyi",
    "İkisi de kabul edilebilir",
    "İkisi de sorunlu",
    "Bağlam olmadan değerlendirilemez",
    "Testten çıkar",
}
READY_DECISIONS = {"4o daha iyi", "Luna daha iyi", "İkisi de kabul edilebilir"}
SPLIT_DECISIONS = {"Bölünmedi", "Gerekli", "Gereksiz tekrar", "Yanlış bölme"}
REQUIRED_COLUMNS = {
    "Vaka",
    "Öğrenci mesajı",
    "Mevcut gold soru",
    "Gold QnA ID",
    "Veri sorunu",
    "Split sayısı",
    "Split kararı",
    "İnsan kararı",
    "Kabul edilen QnA ID'leri",
    "İnceleme notu",
}
INTEGER = re.compile(r"^[0-9]+$")


def file_digest(path: Path) -> str:
    digest = hashlib.blake2b(digest_size=32)
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    return str(value or "").strip()


def _integer(value: Any, *, field: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{field} pozitif tam sayı olmalı")  # noqa: TRY004
    if isinstance(value, int):
        result = value
    elif isinstance(value, float) and value.is_integer():
        result = int(value)
    elif INTEGER.fullmatch(_text(value)):
        result = int(_text(value))
    else:
        raise ValueError(f"{field} pozitif tam sayı olmalı")
    if result < 1:
        raise ValueError(f"{field} pozitif tam sayı olmalı")
    return result


def parse_id_groups(value: Any) -> list[list[int]]:
    raw = _text(value)
    if not raw:
        return []
    groups = []
    for group_index, raw_group in enumerate(raw.split("|"), start=1):
        tokens = [token.strip() for token in raw_group.split(",")]
        if not tokens or any(not token for token in tokens):
            raise ValueError(f"Boş QnA grubu: {group_index}")
        ids = []
        for token in tokens:
            if not INTEGER.fullmatch(token) or int(token) < 1:
                raise ValueError(f"Geçersiz QnA ID: {token!r}")
            ids.append(int(token))
        if len(ids) != len(set(ids)):
            raise ValueError(f"Aynı niyet grubunda tekrarlanan QnA ID: {group_index}")
        groups.append(sorted(ids))
    return groups


def load_catalog(path: Path) -> dict[int, dict[str, Any]]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"QnA kataloğu okunamadı: {path}") from exc
    rows = payload.get("qna") if isinstance(payload, dict) else None
    if not isinstance(rows, list):
        raise TypeError("QnA kataloğunda qna listesi bulunamadı")
    catalog = {}
    for row in rows:
        qna_id = _integer(row.get("id"), field="QnA ID")
        if qna_id in catalog:
            raise ValueError(f"Katalogda tekrarlanan QnA ID: {qna_id}")
        question = _text(row.get("question"))
        answer = _text(row.get("answer"))
        if not question or not answer:
            raise ValueError(f"Katalogda boş soru/cevap: QnA {qna_id}")
        catalog[qna_id] = {
            "id": qna_id,
            "question": question,
            "answer": answer,
        }
    return catalog


def load_baseline_aliases(
    path: Path, catalog: dict[int, dict[str, Any]]
) -> list[dict[str, Any]]:
    by_content: dict[tuple[str, str], list[int]] = {}
    for qna_id, item in catalog.items():
        by_content.setdefault((item["question"], item["answer"]), []).append(qna_id)

    cases = []
    seen_aliases = set()
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"Geçersiz baseline JSONL satırı: {line_number}") from exc
            alias_id = _text(row.get("alias_id"))
            if not alias_id or alias_id in seen_aliases:
                raise ValueError(f"Boş/tekrarlanan alias ID: satır {line_number}")
            seen_aliases.add(alias_id)
            question = _text(row.get("canonical_question"))
            answer = _text(row.get("expected_answer"))
            matches = by_content.get((question, answer), [])
            if len(matches) != 1:
                raise ValueError(
                    f"Baseline QnA aktif katalogda tekil değil: {alias_id}"
                )
            qna_id = matches[0]
            message = _text(row.get("query_text"))
            if not message:
                raise ValueError(f"Baseline öğrenci mesajı boş: {alias_id}")
            cases.append(
                {
                    "case_id": len(cases) + 1,
                    "alias_id": alias_id,
                    "student_message": message,
                    "status": "ready",
                    "reviewed": False,
                    "source_row": None,
                    "review_decision": None,
                    "review_note": "",
                    "data_issue": "",
                    "previous_gold": {"question": question, "qna_id": qna_id},
                    "split_decision": "Bölünmedi",
                    "intents": [
                        {
                            "intent_index": 1,
                            "intent_text": message,
                            "accepted_qna_ids": [qna_id],
                            "accepted_answers": [catalog[qna_id]["answer"]],
                        }
                    ],
                    "replay": {
                        "automatic_replay_eligible": bool(
                            row.get("automatic_replay_eligible")
                        ),
                        "match_status": _text(row.get("match_status")),
                        "session_ids": row.get("matched_session_ids") or [],
                    },
                }
            )
    return cases


def merge_reviewed_cases(
    baseline_cases: list[dict[str, Any]], reviewed_cases: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    merged = [{**case} for case in baseline_cases]
    errors = []
    for review in reviewed_cases:
        case_id = review["case_id"]
        if not 1 <= case_id <= len(merged):
            errors.append(
                {
                    "row": review["source_row"],
                    "case_id": case_id,
                    "field": "Vaka",
                    "code": "review_case_not_in_baseline",
                    "message": "İnceleme vakası 516 alias tabanında bulunamadı",
                }
            )
            continue
        baseline = merged[case_id - 1]
        if baseline["student_message"] != review["student_message"]:
            errors.append(
                {
                    "row": review["source_row"],
                    "case_id": case_id,
                    "field": "Öğrenci mesajı",
                    "code": "review_message_mismatch",
                    "message": "İnceleme mesajı baseline alias mesajıyla eşleşmiyor",
                }
            )
            continue
        merged[case_id - 1] = {
            **baseline,
            **review,
            "alias_id": baseline["alias_id"],
            "reviewed": True,
            "replay": baseline["replay"],
        }
    return merged, errors


def _sheet_rows(sheet, header_row: int) -> list[dict[str, Any]]:
    headers = [_text(cell.value) for cell in sheet[header_row]]
    if len(headers) != len(set(headers)):
        raise ValueError(f"Tekrarlanan sütun başlığı: {sheet.title}")
    rows = []
    for row_number, cells in enumerate(
        sheet.iter_rows(min_row=header_row + 1, max_col=len(headers)),
        start=header_row + 1,
    ):
        values = [cell.value for cell in cells]
        if not any(value is not None and _text(value) for value in values):
            continue
        rows.append({**dict(zip(headers, values, strict=True)), "_row": row_number})
    return rows


def load_review_workbook(
    path: Path,
) -> tuple[list[dict[str, Any]], dict[int, list[str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        required_sheets = {"Karar listesi", "Split vakaları"}
        if not required_sheets.issubset(workbook.sheetnames):
            raise ValueError("Gerekli workbook sekmeleri bulunamadı")
        decisions = _sheet_rows(workbook["Karar listesi"], 4)
        missing = REQUIRED_COLUMNS - set(decisions[0]) if decisions else REQUIRED_COLUMNS
        if missing:
            raise ValueError(f"Eksik karar sütunları: {', '.join(sorted(missing))}")

        split_map = {}
        for row in _sheet_rows(workbook["Split vakaları"], 4):
            case_id = _integer(row.get("Vaka"), field="Vaka")
            if case_id in split_map:
                raise ValueError(f"Tekrarlanan split vakası: {case_id}")
            intents = [
                part.strip() for part in _text(row.get("Luna ayrımı")).split("|")
            ]
            if any(not intent for intent in intents):
                raise ValueError(f"Boş split niyeti: vaka {case_id}")
            split_map[case_id] = intents
        return decisions, split_map
    finally:
        workbook.close()


def apply_simple_review_result(
    rows: list[dict[str, Any]],
    result_path: Path,
    *,
    expected_source_digest: str,
) -> list[dict[str, Any]]:
    """Overlay the browser review export onto workbook review rows."""
    try:
        payload = json.loads(result_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Kolay inceleme sonucu okunamadı: {result_path}") from exc
    if payload.get("schema_version") != "simple-gold-review-result-v1":
        raise ValueError("Kolay inceleme sonucunun şema sürümü geçersiz")
    if payload.get("source_digest") != expected_source_digest:
        raise ValueError("Kolay inceleme sonucu farklı bir workbook'a ait")
    reviews = payload.get("reviews")
    if not isinstance(reviews, list):
        raise TypeError("Kolay inceleme sonucunda reviews listesi bulunamadı")

    by_case: dict[int, dict[str, Any]] = {}
    for review in reviews:
        case_id = _integer(review.get("case_no"), field="Vaka")
        if case_id in by_case:
            raise ValueError(f"Kolay inceleme sonucunda tekrarlanan vaka: {case_id}")
        by_case[case_id] = review

    expected_cases = {_integer(row.get("Vaka"), field="Vaka") for row in rows}
    if set(by_case) != expected_cases:
        missing = sorted(expected_cases - set(by_case))
        extra = sorted(set(by_case) - expected_cases)
        raise ValueError(
            "Kolay inceleme vaka kümesi workbook ile eşleşmiyor: "
            f"eksik={missing}, fazla={extra}"
        )

    overlaid = []
    for row in rows:
        case_id = _integer(row.get("Vaka"), field="Vaka")
        review = by_case[case_id]
        overlaid.append(
            {
                **row,
                "İnsan kararı": _text(review.get("human_decision")),
                "Kabul edilen QnA ID'leri": _text(
                    review.get("accepted_qna_ids")
                ),
                "İnceleme notu": _text(review.get("review_note")),
                "_evaluation_status_override": _text(
                    review.get("evaluation_status_override")
                ),
            }
        )
    return overlaid


def validate_rows(
    rows: list[dict[str, Any]],
    split_map: dict[int, list[str]],
    catalog: dict[int, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    errors = []
    cases = []
    statuses: Counter[str] = Counter()
    seen_cases: set[int] = set()

    for row in rows:
        row_number = int(row.get("_row") or 0)
        row_errors = []
        try:
            case_id = _integer(row.get("Vaka"), field="Vaka")
        except ValueError as exc:
            errors.append(
                {
                    "row": row_number,
                    "field": "Vaka",
                    "code": "invalid_case_id",
                    "message": str(exc),
                }
            )
            continue
        if case_id in seen_cases:
            errors.append(
                {
                    "row": row_number,
                    "case_id": case_id,
                    "field": "Vaka",
                    "code": "duplicate_case_id",
                    "message": "Vaka kimliği tekrarlanıyor",
                }
            )
            continue
        seen_cases.add(case_id)

        def add_error(
            field: str,
            code: str,
            message: str,
            *,
            _row_errors=row_errors,
            _row_number=row_number,
            _case_id=case_id,
        ) -> None:
            _row_errors.append(
                {
                    "row": _row_number,
                    "case_id": _case_id,
                    "field": field,
                    "code": code,
                    "message": message,
                }
            )

        message = _text(row.get("Öğrenci mesajı"))
        if not message:
            add_error("Öğrenci mesajı", "missing_message", "Öğrenci mesajı boş")

        decision = _text(row.get("İnsan kararı"))
        if not decision:
            add_error(
                "İnsan kararı",
                "missing_human_decision",
                "İnsan kararı girilmemiş",
            )
        elif decision not in DECISIONS:
            add_error(
                "İnsan kararı",
                "invalid_human_decision",
                f"Geçersiz karar: {decision}",
            )

        split_decision = _text(row.get("Split kararı"))
        if split_decision not in SPLIT_DECISIONS:
            add_error(
                "Split kararı",
                "invalid_split_decision",
                f"Geçersiz split kararı: {split_decision}",
            )
        try:
            split_count = _integer(row.get("Split sayısı"), field="Split sayısı")
        except ValueError as exc:
            split_count = 0
            add_error("Split sayısı", "invalid_split_count", str(exc))

        try:
            id_groups = parse_id_groups(row.get("Kabul edilen QnA ID'leri"))
        except ValueError as exc:
            id_groups = []
            add_error("Kabul edilen QnA ID'leri", "invalid_qna_ids", str(exc))

        unknown_ids = sorted(
            {
                qna_id
                for group in id_groups
                for qna_id in group
                if qna_id not in catalog
            }
        )
        if unknown_ids:
            add_error(
                "Kabul edilen QnA ID'leri",
                "unknown_qna_ids",
                "Aktif katalogda bulunmayan QnA ID: "
                + ", ".join(map(str, unknown_ids)),
            )

        note = _text(row.get("İnceleme notu"))
        status_override = _text(row.get("_evaluation_status_override"))
        status = "invalid"
        if status_override:
            if status_override not in {"needs_context", "needs_kb", "excluded"}:
                add_error(
                    "Teknik durum",
                    "invalid_status_override",
                    f"Geçersiz teknik durum: {status_override}",
                )
            else:
                status = status_override
                if not note:
                    add_error(
                        "İnceleme notu",
                        "missing_review_note",
                        "Teknik durum değişikliği için açıklama gerekli",
                    )
                if id_groups:
                    add_error(
                        "Kabul edilen QnA ID'leri",
                        "qna_ids_not_allowed",
                        f"{status_override} durumunda QnA ID girilmemeli",
                    )
        elif decision in READY_DECISIONS:
            status = "ready"
            if not id_groups:
                add_error(
                    "Kabul edilen QnA ID'leri",
                    "missing_qna_ids",
                    "Koşulabilir vaka için en az bir QnA ID gerekli",
                )
        elif decision == "İkisi de sorunlu":
            status = "ready" if id_groups else "needs_kb"
            if not note:
                add_error(
                    "İnceleme notu",
                    "missing_review_note",
                    "Sorunlu cevap kararı için açıklama gerekli",
                )
        elif decision == "Bağlam olmadan değerlendirilemez":
            status = "needs_context"
            if not note:
                add_error(
                    "İnceleme notu",
                    "missing_review_note",
                    "Eksik bağlamın ne olduğu yazılmalı",
                )
        elif decision == "Testten çıkar":
            status = "excluded"
            if not note:
                add_error(
                    "İnceleme notu",
                    "missing_review_note",
                    "Testten çıkarma gerekçesi yazılmalı",
                )

        if not status_override and status in {"needs_context", "excluded"} and id_groups:
            add_error(
                "Kabul edilen QnA ID'leri",
                "qna_ids_not_allowed",
                f"{status} durumunda QnA ID girilmemeli",
            )

        if split_decision == "Gerekli" and status == "ready":
            intents = split_map.get(case_id) or []
            if len(intents) != split_count:
                add_error(
                    "Split sayısı",
                    "split_source_mismatch",
                    "Split metinleri ile split sayısı eşleşmiyor",
                )
            if len(id_groups) != split_count:
                add_error(
                    "Kabul edilen QnA ID'leri",
                    "intent_group_count_mismatch",
                    f"{split_count} niyet için {split_count} adet '|' ile "
                    "ayrılmış ID grubu gerekli",
                )
        else:
            intents = [message] if message else []
            if len(id_groups) > 1:
                add_error(
                    "Kabul edilen QnA ID'leri",
                    "unexpected_intent_groups",
                    "Split gerekli değilken birden fazla niyet grubu girilmiş",
                )

        if row_errors:
            status = "invalid"
            errors.extend(row_errors)
        statuses[status] += 1

        intent_records = []
        if status == "ready":
            for intent_index, intent_text in enumerate(intents):
                accepted_ids = (
                    id_groups[intent_index] if intent_index < len(id_groups) else []
                )
                intent_records.append(
                    {
                        "intent_index": intent_index + 1,
                        "intent_text": intent_text,
                        "accepted_qna_ids": accepted_ids,
                        "accepted_answers": [
                            catalog[qna_id]["answer"]
                            for qna_id in accepted_ids
                            if qna_id in catalog
                        ],
                    }
                )

        cases.append(
            {
                "case_id": case_id,
                "source_row": row_number,
                "student_message": message,
                "status": status,
                "review_decision": decision,
                "review_note": note,
                "evaluation_status_override": status_override or None,
                "data_issue": _text(row.get("Veri sorunu")),
                "previous_gold": {
                    "question": _text(row.get("Mevcut gold soru")),
                    "qna_id": _text(row.get("Gold QnA ID")),
                },
                "split_decision": split_decision,
                "intents": intent_records,
            }
        )

    return cases, errors, dict(sorted(statuses.items()))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--qna-catalog", type=Path, required=True)
    parser.add_argument("--baseline-aliases", type=Path, required=True)
    parser.add_argument(
        "--review-json",
        type=Path,
        help="Kolay inceleme ekranından indirilen sonuç dosyası",
    )
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    catalog = load_catalog(args.qna_catalog)
    baseline_cases = load_baseline_aliases(args.baseline_aliases, catalog)
    rows, split_map = load_review_workbook(args.workbook)
    if args.review_json:
        rows = apply_simple_review_result(
            rows,
            args.review_json,
            expected_source_digest=file_digest(args.workbook),
        )
    reviewed_cases, errors, review_status_counts = validate_rows(
        rows, split_map, catalog
    )
    cases, alignment_errors = merge_reviewed_cases(baseline_cases, reviewed_cases)
    errors.extend(alignment_errors)
    report = {
        "schema_version": "gold-v2-validation-v1",
        "valid": not errors,
        "workbook": {
            "name": args.workbook.name,
            "blake2b": file_digest(args.workbook),
        },
        "review_json": (
            {
                "name": args.review_json.name,
                "blake2b": file_digest(args.review_json),
            }
            if args.review_json
            else None
        ),
        "qna_catalog": {
            "name": args.qna_catalog.name,
            "blake2b": file_digest(args.qna_catalog),
            "active_qna_count": len(catalog),
        },
        "baseline_aliases": {
            "name": args.baseline_aliases.name,
            "blake2b": file_digest(args.baseline_aliases),
            "case_count": len(baseline_cases),
        },
        "case_count": len(baseline_cases),
        "review_case_count": len(rows),
        "review_status_counts": review_status_counts,
        "error_count": len(errors),
        "errors": errors,
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    if errors:
        summary_keys = (
            "valid",
            "case_count",
            "review_case_count",
            "review_status_counts",
            "error_count",
        )
        print(
            json.dumps(
                {key: report[key] for key in summary_keys},
                ensure_ascii=False,
                indent=2,
            )
        )
        return 2

    referenced_ids = sorted(
        {
            qna_id
            for case in cases
            for intent in case["intents"]
            for qna_id in intent["accepted_qna_ids"]
        }
    )
    status_counts = dict(sorted(Counter(case["status"] for case in cases).items()))
    gold = {
        "schema_version": "gold-v2",
        "source": report["workbook"],
        "qna_catalog": report["qna_catalog"],
        "status_counts": status_counts,
        "qna_snapshot": [catalog[qna_id] for qna_id in referenced_ids],
        "cases": cases,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(gold, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {"valid": True, "case_count": len(cases), "output": str(args.output)},
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
