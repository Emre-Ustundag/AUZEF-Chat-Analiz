from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.build_session_gold_review_package import (
    CONTEXT_CASES,
    CONTEXT_DECISIONS,
    INTENT_CASES,
    INTENT_DECISIONS,
    PENDING_CASE,
    ROOT,
)

PACKAGE = ROOT / "outputs" / "session-gold-v2-review-20260917"
AUDIT = PACKAGE / "session-gold-human-review.json"
EXCEL = PACKAGE / "session-gold-human-review.xlsx"
pytestmark = pytest.mark.skipif(not AUDIT.exists(), reason="Review paketi üretilmemiş")


@pytest.fixture(scope="module")
def audit() -> dict:
    return json.loads(AUDIT.read_text(encoding="utf-8"))


def test_review_universe_matches_expected_cases(audit):
    universe = audit["review_universe"]
    assert universe["context_cases"] == sorted(CONTEXT_CASES)
    assert universe["intent_cases"] == sorted(INTENT_CASES)
    assert universe["anomaly_cases"] == [345]
    assert universe["total_items"] == len({*CONTEXT_CASES, *INTENT_CASES, *universe["anomaly_cases"]})


def test_pending_case_is_not_in_review_queue(audit):
    assert all(item["case_id"] != PENDING_CASE for item in audit["items"])


def test_cases_appear_once_with_merged_review_types(audit):
    case_ids = [item["case_id"] for item in audit["items"]]
    assert len(case_ids) == len(set(case_ids))
    for item in audit["items"]:
        assert item["review_types"] == sorted(set(item["review_types"]))
        assert item["review_types"], item["case_id"]
    multi = [i["case_id"] for i in audit["items"] if len(i["review_types"]) > 1]
    assert multi == audit["review_universe"]["multi_type_cases"]


def test_frozen_gold_and_session_files_untouched(audit):
    assert audit["frozen_inputs_unchanged"] is True
    for relative, digest in audit["frozen_inputs_sha256_after"].items():
        current = hashlib.sha256((ROOT / relative).read_bytes()).hexdigest()
        assert current == digest, relative


def test_no_decision_is_prefilled_beyond_current_gold(audit):
    for item in audit["items"]:
        assert item["prefilled_final_expected_qna_ids"] == item["current_expected_qna_ids"]
        assert set(item["allowed_decisions"]["context"]) == set(CONTEXT_DECISIONS)
        assert set(item["allowed_decisions"]["intent"]) == set(INTENT_DECISIONS)
        assert "decision" not in item


def test_turns_after_target_are_marked_evidence_only(audit):
    seen = False
    for item in audit["items"]:
        evidence = item.get("evidence_context") or {}
        for turn in evidence.get("turns_after_target_evidence_only", []):
            seen = True
            assert turn["evidence_only_not_evaluator_context"] is True
        for turn in evidence.get("turns_before_target_in_segment", []):
            assert turn["evidence_only_not_evaluator_context"] is False
    assert seen, "hedef sonrası kanıt hiç toplanmamış"


TEXT_KEYS = {"text", "gold_user_message", "gold_message", "gold_message_repr", "target_text",
             "real_text_repr", "why_context_required", "gold_notes", "splitter_note"}


def collect_texts(node, found=None):
    """Yalnız mesaj metinleri; digest alanları taranmaz (hex dizileri yanlış alarm verir)."""
    found = [] if found is None else found
    if isinstance(node, dict):
        for key, value in node.items():
            if key in TEXT_KEYS and isinstance(value, str):
                found.append(value)
            else:
                collect_texts(value, found)
    elif isinstance(node, list):
        for value in node:
            collect_texts(value, found)
    return found


def test_evidence_is_pii_masked(audit):
    texts = collect_texts(audit)
    assert len(texts) > 50
    for pattern in (r"[\w.+-]+@[\w-]+\.[\w.]+", r"(?<!\d)(?:\+90|0)?5\d{9}(?!\d)", r"(?<!\d)[1-9]\d{10}(?!\d)"):
        hits = [t for t in texts if re.search(pattern, t)]
        assert not hits, (pattern, hits[:2])


def test_excel_case_ids_match_json(audit):
    workbook = load_workbook(EXCEL, data_only=False)
    try:
        assert workbook.sheetnames == ["Başlangıç", "1-Context Review", "2-Intent Review", "3-Data Anomaly", "Kaynak Kanıt"]
        context = [row[0] for row in workbook["1-Context Review"].iter_rows(min_row=2, values_only=True) if row[0]]
        intent = [row[0] for row in workbook["2-Intent Review"].iter_rows(min_row=2, values_only=True) if row[0]]
        anomaly = [row[0] for row in workbook["3-Data Anomaly"].iter_rows(min_row=2, values_only=True) if row[0]]
        assert sorted(context) == sorted(CONTEXT_CASES)
        assert sorted(intent) == sorted(INTENT_CASES)
        assert sorted(anomaly) == audit["review_universe"]["anomaly_cases"]
        evidence_cases = {row[0] for row in workbook["Kaynak Kanıt"].iter_rows(min_row=2, values_only=True) if row[0]}
        assert evidence_cases <= {item["case_id"] for item in audit["items"]}
    finally:
        workbook.close()


def test_excel_decision_cells_are_empty_with_validation():
    workbook = load_workbook(EXCEL)
    try:
        sheet = workbook["1-Context Review"]
        assert [cell.value for cell in sheet["C"][1:]] == [None] * len(CONTEXT_CASES)
        assert any("KEEP_CONTEXT_REQUIRED" in (validation.formula1 or "") for validation in sheet.data_validations.dataValidation)
        intent_sheet = workbook["2-Intent Review"]
        assert [cell.value for cell in intent_sheet["B"][1:]] == [None] * len(INTENT_CASES)
        assert any("MARK_MULTI_INTENT" in (validation.formula1 or "")
                   for validation in intent_sheet.data_validations.dataValidation)
    finally:
        workbook.close()


def test_unresolved_context_items_carry_source_evidence(audit):
    items = {item["case_id"]: item for item in audit["items"]}
    for case_id in CONTEXT_CASES:
        evidence = items[case_id]["evidence_context"]
        assert evidence["review_reason"] in {"INSUFFICIENT_REAL_CONTEXT", "SOURCE_NOT_FOUND"}
        assert evidence["why_context_required"]
        if evidence["review_reason"] == "SOURCE_NOT_FOUND":
            assert evidence["exact_match_found"] is False
            assert evidence["searched_sources"]
        else:
            assert evidence["session_id"] and evidence["target_turn_id"]
            assert evidence["raw_messages_before_target_total"] >= 0


def test_intent_items_expose_splitter_pieces(audit):
    items = {item["case_id"]: item for item in audit["items"]}
    for case_id in INTENT_CASES:
        evidence = items[case_id]["evidence_intent"]
        assert len(evidence["intent_pieces"]) >= 2, case_id
        assert evidence["splitter_label"] == "Gerekli"
        for piece in evidence["intent_pieces"]:
            for suggestion in piece["kb_suggestions_heuristic_only"]:
                assert suggestion["ref"].startswith("QNA-")


def test_anomaly_item_states_verified_cause(audit):
    anomaly = next(item for item in audit["items"] if "DATA_ANOMALY_MERGED_MESSAGE" in item["review_types"])
    evidence = anomaly["evidence_anomaly"]
    assert evidence["texts_differ_only_by_case"] is True
    assert "U+0307" in evidence["verified_cause"]
    assert evidence["suggested_fix"]
