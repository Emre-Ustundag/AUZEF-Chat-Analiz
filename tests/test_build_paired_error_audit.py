from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.build_paired_error_audit import BENCH, build, fine_cause, model_diff_labels

OUTPUT = Path("outputs/paired-error-audit-v1-20260918")
QNA = {1: {"question_text": "Ders kaydı nasıl yapılır?", "answer_text": "Kayıt ekranından."},
       2: {"question_text": "Ders kaydı ne zaman yapılır?", "answer_text": "Takvimde ilan edilir."},
       3: {"question_text": "Harç iadesi", "answer_text": "Başvuru ile."}}


def view(**overrides):
    base = {"exact_correct": False, "relaxed_correct": False, "split_count": 1, "used_asks": [],
            "speculative_correct": False, "speculative_gold_rank": {"qna_rank": 1}, "fallback": False, "source": "llm",
            "calendar_selected": False, "declined_any": False, "selected_qna_ids": [], "raw_select_outputs": ["1"],
            "speculative_candidates": [None, 1, 2], "speculative_declined": False}
    return {**base, **overrides}


def ask(selected, gold_rank=1, selected_rank=None, declined=False, calendar=False):
    return {"selected_qna_id": selected, "declined": declined, "selected_calendar": calendar,
            "gold_rank": {"qna_rank": gold_rank, "selected_qna_rank": selected_rank}}


CASE = {"expected_qna_ids": [1], "expected_intent_groups": [[1]], "turn_type": "FIRST_TURN", "user_message": "mesaj"}


@pytest.mark.parametrize(
    ("model_view", "expected"),
    [
        (view(exact_correct=True), None),
        (view(relaxed_correct=True, split_count=2, used_asks=[ask(1), ask(3)], selected_qna_ids=[1, 3]), "FALSE_SPLIT_SINGLE_INTENT"),
        (view(split_count=2, speculative_correct=True, used_asks=[ask(3, gold_rank=None), ask(2, gold_rank=None)],
              selected_qna_ids=[2, 3]), "FRAGMENT_LOST_CONTEXT"),
        (view(used_asks=[ask(3, gold_rank=None)], selected_qna_ids=[3]), "RETRIEVAL_TRUE_MISS"),
        (view(used_asks=[ask(None, declined=True)], declined_any=True, fallback=True), "SELECTOR_NONE_WHEN_GOLD_PRESENT"),
        (view(used_asks=[ask(None, calendar=True)], calendar_selected=True), "SELECTOR_CALENDAR_DISTRACTION"),
        (view(used_asks=[ask(2, selected_rank=2)], selected_qna_ids=[2]), "SELECTOR_NEAR_DUPLICATE_CONFUSION"),
        (view(used_asks=[ask(3, selected_rank=2)], selected_qna_ids=[3]), "SELECTOR_WRONG_SEMANTIC_MATCH"),
        (view(raw_select_outputs=["bilmiyorum"], used_asks=[ask(None)]), "FORMAT_OR_PARSE_ERROR"),
    ],
)
def test_fine_cause_rules(model_view, expected):
    cause, _ = fine_cause(CASE, model_view, QNA, {}, view())
    assert cause == expected


def test_kb_overlap_when_message_is_alias_of_selected():
    cause, secondary = fine_cause(CASE, view(used_asks=[ask(3, selected_rank=2)], selected_qna_ids=[3]), QNA, {"mesaj": {3}}, view())
    assert cause == "KB_OVERLAP_AMBIGUITY" and "GOLD_REVIEW_CANDIDATE" in secondary


def test_missed_multi_intent():
    case = {**CASE, "expected_qna_ids": [1, 3], "expected_intent_groups": [[1], [3]]}
    cause, _ = fine_cause(case, view(used_asks=[ask(1)], selected_qna_ids=[1]), QNA, {}, view())
    assert cause == "MISSED_MULTI_INTENT"


def test_pure_selector_difference_requires_same_input_and_split():
    right = view(exact_correct=True, selected_qna_ids=[1])
    wrong = view(selected_qna_ids=[2])
    assert "PURE_SELECTOR_MODEL_DIFFERENCE" in model_diff_labels(right, wrong)
    split = view(split_count=2, selected_qna_ids=[2])
    labels = model_diff_labels(right, split)
    assert "PURE_SELECTOR_MODEL_DIFFERENCE" not in labels and "MODEL_DIFF_FRAGMENTATION" in labels


@pytest.mark.skipif(not (BENCH / "4o-mini" / "results.jsonl").exists(), reason="benchmark izi yok")
class TestAuditOutputs:
    @pytest.fixture(scope="class")
    def built(self, tmp_path_factory):
        before = {p.name: hashlib.new("sha256", p.read_bytes()).hexdigest() for p in BENCH.rglob("*") if p.is_file()}
        result = build(tmp_path_factory.mktemp("audit"))
        after = {p.name: hashlib.new("sha256", p.read_bytes()).hexdigest() for p in BENCH.rglob("*") if p.is_file()}
        assert before == after, "benchmark çıktıları değişmemeli"
        return result

    def test_universe_matches_exact_metric_buckets(self, built):
        universe = built["report"]["universe"]
        comparison = json.loads((BENCH / "comparison.json").read_text(encoding="utf-8"))["paired"]["exact_metric"]["counts"]
        assert universe["only_4o_mini_correct"] == comparison["only_4o_mini"]
        assert universe["only_luna_high_correct"] == comparison["only_luna_high"]
        assert universe["both_wrong"] == comparison["both_wrong"]

    def test_every_failure_has_single_primary_cause(self, built):
        for case in built["cases"]:
            for key in ("4o_mini", "luna_high"):
                view_ = case[key]
                assert (view_["fine_cause"] is None) == bool(view_["exact_correct"])

    def test_multi_intent_loss_stages(self, built):
        cases = {c["case_id"]: c for c in built["cases"]}
        assert cases[480]["4o_mini"]["fine_cause"] == "SELECTOR_NONE_WHEN_GOLD_PRESENT"
        assert cases[480]["luna_high"]["fine_cause"] == "MISSED_MULTI_INTENT"
        multi = built["report"]["multi_intent"]
        assert multi["480"]["4o-mini"]["lost_groups"][0]["loss_stage"] == "SELECTOR"
        assert multi["480"]["luna-high"]["lost_groups"][0]["loss_stage"] == "SPLITTER"

    def test_rebuild_is_deterministic(self, built, tmp_path):
        second = build(tmp_path)
        assert second["cases"] == built["cases"]
        assert second["report"]["root_causes"] == built["report"]["root_causes"]

    @pytest.mark.skipif(not OUTPUT.exists(), reason="denetim çıktısı yok")
    def test_review_workbook_matches_json(self, built):
        workbook = load_workbook(OUTPUT / "paired-error-human-review.xlsx")
        try:
            expected = {"Özet", "Only 4o Correct", "Only Luna Correct", "Both Wrong", "False Splits", "Calendar",
                        "Multi Intent", "Needs Human Review"}
            assert expected <= set(workbook.sheetnames)
            universe = built["report"]["universe"]
            for sheet, count in (("Only 4o Correct", universe["only_4o_mini_correct"]),
                                 ("Only Luna Correct", universe["only_luna_high_correct"]),
                                 ("Both Wrong", universe["both_wrong"]),
                                 ("Needs Human Review", built["report"]["human_review"]["cases"])):
                rows = [r for r in workbook[sheet].iter_rows(min_row=2, values_only=True) if r[0]]
                assert len(rows) == count, sheet
        finally:
            workbook.close()
